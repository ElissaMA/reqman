"""提醒单生成器测试（模板填充/黑红字/重点黄底）"""
import io

import openpyxl

from reqman.services.reminder_generator import generate_reminder

YELLOW = "FFFFFF00"


class TestGenerateReminder:
    def _items(self):
        return [
            {"task_name": "电子例行卡", "category": "电子", "reminder_type": "一般提醒", "source": "例行"},
            {"task_name": "发动机其他卡", "category": "发动机", "reminder_type": "重点提醒", "source": "其他"},
            {"task_name": "机体卡", "category": "机体", "reminder_type": "一般提醒", "source": "例行"},
        ]

    def test_output_structure(self):
        form = {"reg": "B-8323", "aircraft_type": "A320", "description": "46A",
                "date": "2026.08.23", "routine_count": 10, "other_count": 2}
        buffer, filename = generate_reminder(form, self._items())
        assert isinstance(buffer, io.BytesIO)
        assert "定检工作提醒单（B-8323 46A）2026.08.23.xlsx" in filename
        wb = openpyxl.load_workbook(buffer)
        ws = wb["工卡提醒"]
        assert "定检工作提醒单" in str(ws["A1"].value)
        assert "时间：" in str(ws["A2"].value)
        assert "总份数：10+2" in str(ws["C2"].value)
        assert "机号：B-8323" in str(ws["A3"].value)
        # 数据区：电子→A7 例行黑字；发动机→B7 其他红字+黄底；机体→C7
        assert ws["A7"].value == "电子例行卡"
        assert ws["B7"].value == "发动机其他卡"
        assert ws["C7"].value == "机体卡"
        assert ws["A7"].font.color.rgb.endswith("000000")      # 例行黑字
        assert ws["B7"].font.color.rgb.endswith("FF0000")      # 其他红字
        assert ws["B7"].fill.start_color.rgb.endswith(YELLOW)  # 重点黄底
        assert ws["C7"].fill.patternType is None or not ws["C7"].fill.start_color.rgb.endswith(YELLOW)
        wb.close()

    def test_aircraft_info_cells(self):
        """B2=定检描述、C3=APU、A4=FSN、B4=MSN"""
        form = {"reg": "B-100", "aircraft_type": "A320", "description": "46A",
                "level": "46A", "fsn": "F-123", "msn": "MSN-88", "apu": "APU-9",
                "date": "2026.08.23"}
        buffer, _ = generate_reminder(form, [])
        wb = openpyxl.load_workbook(buffer)
        ws = wb["工卡提醒"]
        assert "定检级别：46A" in str(ws["B2"].value)
        assert "APU型号：APU-9" in str(ws["C3"].value)
        assert "FSN：F-123" in str(ws["A4"].value)
        assert "MSN：MSN-88" in str(ws["B4"].value)
        wb.close()

    def test_b2_prefers_description_over_level(self):
        """B2 输出工作包描述（AMRO REVTITLE），描述缺失时回退定检级别。"""
        form = {"reg": "B-100", "description": "A320 4C检", "level": "4C",
                "date": "2026.08.23"}
        buffer, _ = generate_reminder(form, [])
        wb = openpyxl.load_workbook(buffer)
        ws = wb["工卡提醒"]
        assert str(ws["B2"].value) == "定检级别：A320 4C检"
        wb.close()

        form2 = {"reg": "B-100", "description": "", "level": "4C", "date": "2026.08.23"}
        buffer2, _ = generate_reminder(form2, [])
        wb2 = openpyxl.load_workbook(buffer2)
        assert str(wb2["工卡提醒"]["B2"].value) == "定检级别：4C"
        wb2.close()

    def test_second_row_appends(self):
        """同专业多条 → 依次写入 7、8 行"""
        form = {"reg": "B-1", "aircraft_type": "", "description": "", "date": "2026.08.23"}
        buffer, _ = generate_reminder(form, [
            {"task_name": "卡1", "category": "电子", "reminder_type": "一般提醒", "source": "例行"},
            {"task_name": "卡2", "category": "电子", "reminder_type": "重点提醒", "source": "其他"},
        ])
        wb = openpyxl.load_workbook(buffer)
        ws = wb["工卡提醒"]
        assert ws["A7"].value == "卡1"
        assert ws["A8"].value == "卡2"
        wb.close()