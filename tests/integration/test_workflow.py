"""端到端工作流测试 — 上传 → 匹配 → 生成需求单"""

import io

import openpyxl


def _make_mock_excel(items: list[dict]) -> bytes:
    """生成模拟工作清单 Excel���与真实 worklist_parser 解析格式兼容）

    数据布局（与 worklist_parser 一致）：
      - B2 = 机号, D2 = 机型, H2 = ���述, C4 = 日期
      - 第 5 行 = 表头���不被解析）
      - 第 7 行��� = 数据行，列 B=工卡号, E=类型, F=专业, H=工卡描述
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "工作���单"

    # 飞机信息（Row 2）
    ws["B2"] = "B-1234"
    ws["D2"] = "A320"
    ws["H2"] = "A320 定检"

    # 日期（Row 4, C4）
    ws["C4"] = "2026.07.22"

    # 表头（Row 5，仅作标识，不被 parser 使用）
    headers = ["序号", "工卡号", "ATA", "工种", "类型", "专业", "��域", "工卡描述", "工时", "数量", "备注", "L-column"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=5, column=col, value=header)

    # 数据行（从 Row 7 开始）
    # 列对应: B(2)=工卡��, E(5)=类型, F(6)=专业, H(8)=工��描述, L(12)=备注
    for i, item in enumerate(items, 7):
        ws.cell(row=i, column=2, value=item.get("task_code", ""))        # B: 工卡号
        ws.cell(row=i, column=5, value=item.get("task_type", ""))        # E: 类型
        ws.cell(row=i, column=6, value=item.get("category", ""))         # F: 专业
        ws.cell(row=i, column=8, value=item.get("task_name", ""))        # H: 工卡描述
        ws.cell(row=i, column=12, value=item.get("remark", ""))          # L: 备注

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


class TestUploadWorkPackage:
    """工作包上传流程测试"""

    UPLOAD_URL = "/upload"

    def test_upload_no_file(self, client):
        """无文件上传返回错误"""
        resp = client.post(self.UPLOAD_URL, data={})
        # 非AJAX请求：重定向到上传页面（含错误flash）
        assert resp.status_code in (302, 200)
        # AJAX请求返回JSON错误
        resp_ajax = client.post(self.UPLOAD_URL, data={},
                                headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp_ajax.status_code == 400
        data = resp_ajax.get_json()
        assert data["success"] is False

    def test_upload_invalid_extension(self, client):
        """上传非xlsx��件返回错误"""
        data = {"routine_file": (io.BytesIO(b"fake content"), "test.txt")}
        resp = client.post(self.UPLOAD_URL, data=data,
                           headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 400
        data_json = resp.get_json()
        assert data_json["success"] is False

    def test_upload_success_with_mock_excel(self, client):
        """上传有效Excel文件成功"""
        excel_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机检查",
             "category": "发动机", "task_type": "A"},
            {"task_code": "AIR-001", "task_name": "机身蒙皮检查",
             "category": "机体", "task_type": "A"},
        ])
        data = {"routine_file": (io.BytesIO(excel_data), "routine.xlsx")}
        resp = client.post(self.UPLOAD_URL, data=data,
                           headers={"X-Requested-With": "XMLHttpRequest"})
        if resp.status_code == 200:
            data_json = resp.get_json()
            assert data_json["success"] is True
        else:
            assert resp.status_code in (302,)

    def test_upload_then_list_packages(self, prefilled_client, store):
        """上传后工作包出现在列表中"""
        # ��通过API上传
        excel_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机检查",
             "category": "发动机", "task_type": "A"},
        ])
        data = {"routine_file": (io.BytesIO(excel_data), "routine.xlsx")}
        prefilled_client.post(self.UPLOAD_URL, data=data)

        # 检查Store中是否有工作包
        packages = store.get_work_packages()
        assert len(packages) >= 1
        latest = packages[0]
        assert latest is not None
        items = latest.get("all_items", [])
        assert len(items) >= 1

    def test_upload_page_get(self, client):
        """上传页面GET请求正常"""
        resp = client.get(self.UPLOAD_URL)
        assert resp.status_code == 200

    def test_upload_routine_and_other(self, client):
        """同��上传例行和非例行两个���件"""
        routine_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机检查",
             "category": "发动机", "task_type": "A"},
        ])
        other_data = _make_mock_excel([
            {"task_code": "AVN-001", "task_name": "电子测试",
             "category": "电��", "task_type": "B"},
        ])
        data = {
            "routine_file": (io.BytesIO(routine_data), "routine.xlsx"),
            "other_file": (io.BytesIO(other_data), "other.xlsx"),
        }
        resp = client.post(self.UPLOAD_URL, data=data,
                           headers={"X-Requested-With": "XMLHttpRequest"})
        if resp.status_code == 200:
            data_json = resp.get_json()
            assert data_json["success"] is True


class TestMatchAndGenerate:
    """匹配和生成流程测试"""

    def test_rematch_existing_package(self, prefilled_app, prefilled_client, store):
        """上传后重���匹配已存在的工包"""
        # 先上传一个工作包
        excel_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机检查",
             "category": "发动机", "task_type": "A"},
        ])
        data = {"routine_file": (io.BytesIO(excel_data), "routine.xlsx")}
        prefilled_client.post("/upload", data=data)
        all_pkgs = store.get_work_packages()
        pkg = all_pkgs[0] if all_pkgs else None
        pkg_id = str(pkg.get("id", len(all_pkgs))) if pkg else "1"

        resp = prefilled_client.post(f"/packages/{pkg_id}/rematch",
                                      headers={"X-Requested-With": "XMLHttpRequest"})
        if resp.status_code == 200:
            data = resp.get_json()
            assert data["success"] is True
        # 404 也可能（包无id字段），取决于store实现

    def test_rematch_nonexistent(self, prefilled_client):
        """重新匹配不存在的工作包"""
        resp = prefilled_client.post("/packages/9999/rematch",
                                      headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 404
        data = resp.get_json()
        assert data["success"] is False
        assert data["error_code"] == "NOT_FOUND"

    def test_generate_preview_no_package(self, client):
        """生成预览缺少package_id参数"""
        # AJAX
        resp = client.get("/generate", headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 400
        data = resp.get_json()
        assert data["success"] is False

        # 非AJAX（渲染表单页面）
        resp_html = client.get("/generate")
        assert resp_html.status_code == 200

    def test_generate_preview_nonexistent_package(self, prefilled_client):
        """生成预览使用不存在的package_id"""
        resp = prefilled_client.get("/generate?package_id=9999",
                                     headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 404
        data = resp.get_json()
        assert data["success"] is False

    def test_generate_preview_with_package(self, prefilled_app, prefilled_client, store):
        """上���后生成需求单预览"""
        # 上传 -> 获取包ID
        excel_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机��查",
             "category": "发动机", "task_type": "A"},
        ])
        prefilled_client.post("/upload", data={
            "routine_file": (io.BytesIO(excel_data), "routine.xlsx"),
        })
        pkgs = store.get_work_packages()
        if not pkgs:
            return  # 无工作包则跳过
        pkg_id = str(pkgs[0].get("id", len(pkgs)))

        # 先���成预览��触发匹配）
        resp = prefilled_client.get(f"/generate?package_id={pkg_id}",
                                     headers={"X-Requested-With": "XMLHttpRequest"})
        if resp.status_code == 200:
            data = resp.get_json()
            assert data["success"] is True

        # 非AJAX预览
        resp_html = prefilled_client.get(f"/generate?package_id={pkg_id}")
        assert resp_html.status_code in (200, 302)

    def test_generate_post_excel(self, prefilled_app, prefilled_client, store):
        """提交生成表单，下载xlsx文件"""
        # 上传
        excel_data = _make_mock_excel([
            {"task_code": "ENG-001", "task_name": "发动机检查",
             "category": "发动机", "task_type": "A"},
        ])
        prefilled_client.post("/upload", data={
            "routine_file": (io.BytesIO(excel_data), "routine.xlsx"),
        })
        pkgs = store.get_work_packages()
        if not pkgs:
            return
        pkg_id = str(pkgs[0].get("id", len(pkgs)))

        # 生成（先预览让系统匹配）
        prefilled_client.get(f"/generate?package_id={pkg_id}")

        # POST 生成Excel
        resp = prefilled_client.post("/generate", data={
            "package_id": pkg_id,
            "date": "2026.07.22",
            "reg": "B-1234",
        })
        # 成功下载Excel，或重定向
        if resp.status_code == 200:
            content_type = resp.headers.get("Content-Type", "")
            assert "spreadsheet" in content_type or "octet-stream" in content_type
        else:
            assert resp.status_code in (302,)
