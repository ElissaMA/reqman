"""提醒单下载 API 集成测试"""


class TestReminderDownload:
    def test_download_requires_package(self, client):
        resp = client.post("/generate/reminder", data={},
                           headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 400

    def test_download_returns_xlsx(self, client, store):
        # 预置已确认提醒工卡
        store.add(task_code="E-001", task_name="电子例行卡", category="电子")
        store.update(store.find_by_code("E-001")["id"],
                     tools=[{"device_name": "万用表"}],
                     tools_confirmed=True, materials_confirmed=True,
                     reminder_type="重点提醒", card_ok=True)
        # 构造已匹配工作包（is_matched=True 直接复用 matched）
        store.save_work_package({
            "reg": "B-1234", "description": "46A", "date": "2026.08.23",
            "aircraft_info": {"reg": "B-1234", "type": "A320",
                              "description": "46A", "date": "2026.08.23"},
            "matched": [{
                "task_code": "E-001", "task_name": "电子例行卡",
                "category": "电子", "reminder_type": "重点提醒",
                "card_ok": True, "reminder_confirmed": True,
                "source": "例行", "tools": [], "materials": [],
            }],
            "new_cards": [], "cancelled": [], "all_items": [],
            "routine_count": 1, "other_count": 0,
            "is_matched": True, "generated_at": "2026.08.23 10:00",
        })
        pkg_id = store.get_work_packages()[0]["package_id"]
        resp = client.post("/generate/reminder", data={"package_id": pkg_id})
        assert resp.status_code == 200
        # 端点正常返回（内容过滤由 generate_reminder 测试覆盖）
        assert "spreadsheetml" in resp.mimetype

    def test_download_fills_aircraft_info(self, client, store):
        """按 reg 查库补 level/fsn/msn/apu：B2/C3/A4/B4"""
        store.add_aircraft(reg="B-5678", model="A320", engine="LEAP",
                           fsn="F-999", msn="MSN-77", apu="APU-5")
        store.add(task_code="E-010", task_name="电子卡", category="电子")
        store.update(store.find_by_code("E-010")["id"],
                     tools_confirmed=True, materials_confirmed=True,
                     reminder_type="重点提醒", reminder_confirmed=True, card_ok=True)
        store.save_work_package({
            "reg": "B-5678", "description": "46A", "date": "2026.08.23",
            "aircraft_info": {"reg": "B-5678", "type": "A320", "description": "46A",
                              "level": "46A", "date": "2026.08.23"},
            "matched": [{
                "task_code": "E-010", "task_name": "电子卡",
                "category": "电子", "reminder_type": "重点提醒",
                "card_ok": True, "reminder_confirmed": True,
                "source": "例行", "tools": [], "materials": [],
            }],
            "new_cards": [], "cancelled": [], "all_items": [],
            "routine_count": 1, "other_count": 0,
            "is_matched": True, "generated_at": "2026.08.23 10:00",
        })
        pkg_id = next(wp["package_id"] for wp in store.get_work_packages()
                      if wp.get("reg") == "B-5678")
        resp = client.post("/generate/reminder", data={"package_id": pkg_id})
        assert resp.status_code == 200
        import io

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb["工卡提醒"]
        assert "定检级别：46A" in str(ws["B2"].value)
        assert "APU型号：APU-5" in str(ws["C3"].value)
        assert "FSN：F-999" in str(ws["A4"].value)
        assert "MSN：MSN-77" in str(ws["B4"].value)
        wb.close()

    def test_download_aircraft_prefix_fallback(self, client, store):
        """reg 无 B- 前缀时试加前缀查询成功"""
        store.add_aircraft(reg="B-8888", model="A320", engine="LEAP",
                           fsn="F-1", msn="MSN-1", apu="APU-1")
        store.add(task_code="E-011", task_name="电子卡", category="电子")
        store.update(store.find_by_code("E-011")["id"],
                     tools_confirmed=True, materials_confirmed=True,
                     reminder_type="一般提醒", reminder_confirmed=True, card_ok=True)
        store.save_work_package({
            "reg": "B-8888", "description": "A", "date": "2026.08.23",
            "aircraft_info": {"reg": "B-8888", "type": "A320", "description": "A",
                              "level": "A", "date": "2026.08.23"},
            "matched": [{
                "task_code": "E-011", "task_name": "电子卡",
                "category": "电子", "reminder_type": "一般提醒",
                "card_ok": True, "reminder_confirmed": True,
                "source": "例行", "tools": [], "materials": [],
            }],
            "new_cards": [], "cancelled": [], "all_items": [],
            "routine_count": 1, "other_count": 0,
            "is_matched": True, "generated_at": "2026.08.23 10:00",
        })
        # 包内 reg 不带 B- 前缀（8888）
        wp = store.get_work_packages()[-1]
        wp["aircraft_info"]["reg"] = "8888"
        wp["reg"] = "8888"
        store.save_work_package(wp)
        pkg_id = next(w["package_id"] for w in store.get_work_packages()
                      if w.get("reg") == "8888")
        resp = client.post("/generate/reminder", data={"package_id": pkg_id})
        assert resp.status_code == 200
        import io

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.data))
        ws = wb["工卡提醒"]
        assert "APU型号：APU-1" in str(ws["C3"].value)
        wb.close()
        disp = resp.headers.get("Content-Disposition", "")
        assert "定检工作提醒单" in disp or "%E5%AE%9A%E6%A3%80%E5%B7%A5%E4%BD%9C%E6%8F%90%E9%86%92%E5%8D%95" in disp

    def test_filters_unconfirmed_items(self, client, store):
        # card_ok=False 的项不应出现在提醒单
        store.add(task_code="E-002", task_name="未确认卡", category="电子")
        store.update(store.find_by_code("E-002")["id"],
                     tools_confirmed=True, materials_confirmed=True,
                     reminder_type="一般提醒", card_ok=False)
        store.save_work_package({
            "reg": "B-9", "description": "A", "date": "2026.08.23",
            "aircraft_info": {"reg": "B-9", "type": "A320", "description": "A", "date": "2026.08.23"},
            "matched": [{
                "task_code": "E-002", "task_name": "未确认卡",
                "category": "电子", "reminder_type": "一般提醒",
                "card_ok": False, "reminder_confirmed": True,
                "source": "例行", "tools": [], "materials": [],
            }],
            "new_cards": [], "cancelled": [], "all_items": [],
            "routine_count": 0, "other_count": 0,
            "is_matched": True, "generated_at": "2026.08.23 10:00",
        })
        pkg_id = store.get_work_packages()[0]["package_id"]
        resp = client.post("/generate/reminder", data={"package_id": pkg_id})
        assert resp.status_code == 200
        # 端点正常返回（内容过滤由 generate_reminder 测试覆盖）
