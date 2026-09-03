"""库存预警数据库：增改查路由 + 阈值上色 + 查询回写缓存库存。"""

from pathlib import Path

import openpyxl

from reqman.models.json_store import JsonStore
from reqman.services.connectors.session import LoginSessionStore
from reqman.services.inventory_service import InventoryService, _is_warning
from reqman.services.xlsx_workbook import read_demand, write_inventory_copy

RED = "FF0000"
YELLOW = "FFFF00"


# ============================================================
# 路由：新增 / 编辑 / 删除（弹窗表单 + AJAX）
# ============================================================

class TestWarningCRUD:
    def test_new_form_page(self, client):
        resp = client.get("/inventory-warning/new")
        assert resp.status_code == 200
        assert "新增库存预警" in resp.get_data(as_text=True)

    def test_create_ajax_success(self, client, store, ajax_headers):
        resp = client.post(
            "/inventory-warning/new",
            data={"part_number": "km-test-a", "name": "螺钉", "threshold": "10", "note": "常用"},
            headers=ajax_headers,
        )
        data = resp.get_json()
        assert data["success"] is True
        w = store.get_inventory_warning("KM-TEST-A")
        assert w is not None
        assert w["name"] == "螺钉"
        assert w["threshold"] == 10.0
        assert w["stock"] is None
        assert w["note"] == "常用"
        # 清理
        store.delete_inventory_warning("KM-TEST-A")

    def test_create_missing_part_number(self, client, store, ajax_headers):
        resp = client.post(
            "/inventory-warning/new",
            data={"part_number": "   ", "threshold": "10"},
            headers=ajax_headers,
        )
        data = resp.get_json()
        assert data["success"] is False
        assert "件号" in data["message"]

    def test_create_invalid_threshold(self, client, store, ajax_headers):
        resp = client.post(
            "/inventory-warning/new",
            data={"part_number": "KM-TEST-B", "threshold": "abc"},
            headers=ajax_headers,
        )
        data = resp.get_json()
        assert data["success"] is False
        assert "警戒线" in data["message"]

    def test_create_negative_threshold(self, client, store, ajax_headers):
        resp = client.post(
            "/inventory-warning/new",
            data={"part_number": "KM-TEST-C", "threshold": "-5"},
            headers=ajax_headers,
        )
        data = resp.get_json()
        assert data["success"] is False

    def test_edit_ajax_success(self, client, store, ajax_headers):
        store.save_inventory_warning(
            {"part_number": "KM-TEST-D", "name": "旧名", "threshold": 5.0, "note": "旧备注"}
        )
        try:
            resp = client.post(
                "/inventory-warning/KM-TEST-D/edit",
                data={"name": "新名", "threshold": "20", "note": "新备注"},
                headers=ajax_headers,
            )
            data = resp.get_json()
            assert data["success"] is True
            w = store.get_inventory_warning("KM-TEST-D")
            assert w["name"] == "新名"
            assert w["threshold"] == 20.0
            assert w["note"] == "新备注"
            assert w["stock"] is None  # 表单不覆盖缓存库存
        finally:
            store.delete_inventory_warning("KM-TEST-D")

    def test_edit_preserves_threshold_when_omitted(self, client, store, ajax_headers):
        store.save_inventory_warning(
            {"part_number": "KM-TEST-E", "name": "旧名", "threshold": 7.0}
        )
        try:
            resp = client.post(
                "/inventory-warning/KM-TEST-E/edit",
                data={"name": "改名"},  # 不传 threshold
                headers=ajax_headers,
            )
            assert resp.get_json()["success"] is True
            w = store.get_inventory_warning("KM-TEST-E")
            assert w["threshold"] == 7.0  # 原警戒线保留
        finally:
            store.delete_inventory_warning("KM-TEST-E")

    def test_edit_missing_returns_404(self, client, store, ajax_headers):
        resp = client.post(
            "/inventory-warning/KM-NOPE/edit",
            data={"name": "x"},
            headers=ajax_headers,
        )
        assert resp.status_code == 404
        assert resp.get_json()["success"] is False

    def test_delete_ajax(self, client, store, ajax_headers):
        store.save_inventory_warning({"part_number": "KM-TEST-F", "threshold": 1.0})
        assert store.get_inventory_warning("KM-TEST-F") is not None
        resp = client.post("/inventory-warning/KM-TEST-F/delete", headers=ajax_headers)
        data = resp.get_json()
        assert data["success"] is True
        assert store.get_inventory_warning("KM-TEST-F") is None

    def test_delete_url_encoded_part_number(self, client, store, ajax_headers):
        """含空格件号经 urlencode 后仍能正确路由并删除（验证 %20 解码还原）。"""
        store.save_inventory_warning({"part_number": "KM TEST G", "threshold": 1.0})
        assert store.get_inventory_warning("KM TEST G") is not None
        # 模板渲染为 /inventory-warning/KM%20TEST%20G/delete
        resp = client.post("/inventory-warning/KM%20TEST%20G/delete", headers=ajax_headers)
        data = resp.get_json()
        assert data["success"] is True
        assert store.get_inventory_warning("KM TEST G") is None

    def test_delete_missing_returns_404(self, client, ajax_headers):
        resp = client.post("/inventory-warning/KM-NOPE/delete", headers=ajax_headers)
        assert resp.status_code == 404


# ============================================================
# 页面渲染：预警面板 + 查询提示框
# ============================================================

class TestWarningIndex:
    def test_index_lists_warnings_and_summary_box(self, client, store):
        store.save_inventory_warning(
            {"part_number": "KM-PANEL", "name": "垫圈", "threshold": 3.0, "stock": 1.0}
        )
        try:
            resp = client.get("/inventory")
            assert resp.status_code == 200
            html = resp.get_data(as_text=True)
            assert "库存预警数据库" in html
            assert "KM-PANEL" in html
            # 整行着色：库存 1 < 阈值 3 → 标黄
            assert "table-warning" in html
            # 查询提示框容器存在
            assert "amro_last_query" in html or "查询库存进行中" in html or "上次查询" in html
        finally:
            store.delete_inventory_warning("KM-PANEL")

    def test_index_out_of_stock_whole_row_red(self, client, store):
        store.save_inventory_warning(
            {"part_number": "KM-RED", "name": "缺货件", "threshold": 5.0, "stock": 0.0}
        )
        try:
            resp = client.get("/inventory")
            html = resp.get_data(as_text=True)
            assert "KM-RED" in html
            # 库存为 0 → 整行标红（缺货）
            assert "table-danger" in html
        finally:
            store.delete_inventory_warning("KM-RED")


# ============================================================
# 单元：_is_warning 判定
# ============================================================

class TestIsWarning:
    def test_in_db_below_threshold(self):
        assert _is_warning("PN", 2, 4, {"PN": 5}) is True   # 4 < 5

    def test_in_db_at_threshold(self):
        assert _is_warning("PN", 2, 5, {"PN": 5}) is False  # 5 不 < 5

    def test_not_in_db_fallback_yellow(self):
        assert _is_warning("PN", 2, 3, None) is True        # 2 <= 3 < 4

    def test_not_in_db_below_qty_no_warning(self):
        assert _is_warning("PN", 2, 1, None) is False        # 1 < 2 属短缺不标黄

    def test_not_in_db_above_qty_plus_2(self):
        assert _is_warning("PN", 2, 5, None) is False        # 5 >= 4


# ============================================================
# 单元：write_inventory_copy 阈值上色
# ============================================================

def _build_threshold_demand(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "需求单"
    ws["A14"] = "定检专业\n（航材）"
    # PN-001 在库内（阈值5）；PN-002/003 库外；PN-004 缺货
    ws["A15"] = "发动机"; ws["B15"] = "件A"; ws["C15"] = "PN-001"; ws["E15"] = "2"
    ws["A16"] = "发动机"; ws["B16"] = "件B"; ws["C16"] = "PN-002"; ws["E16"] = "5"
    ws["A17"] = "发动机"; ws["B17"] = "件C"; ws["C17"] = "PN-003"; ws["E17"] = "1"
    ws["A18"] = "发动机"; ws["B18"] = "件D"; ws["C18"] = "PN-004"; ws["E18"] = "5"
    wb.save(path)


class TestWriteThreshold:
    def test_threshold_yellow_fallback_off(self, tmp_path: Path):
        """库内件号按阈值标黄，即便按旧规则（库存<使用量+2）不标黄。"""
        src = tmp_path / "demand.xlsx"
        _build_threshold_demand(src)
        rows = read_demand(src)
        pn_cells, pn_qty = {}, {}
        for r in rows:
            pn_cells.setdefault(r.part_number, []).append(r.stock_cell)
            pn_qty.setdefault(r.part_number, []).append(r.qty)
        results = {"PN-001": 4.0, "PN-002": 6.0, "PN-003": 5.0, "PN-004": 1.0}
        # PN-001 库内阈值 5：库存4 < 5 → 黄；旧规则 2<=4<4 假 → 不黄（验证阈值路径）
        buf, _ = write_inventory_copy(
            src, results, pn_qty, pn_cells,
            timestamp_suffix="T", warning_thresholds={"PN-001": 5.0},
        )
        wb = openpyxl.load_workbook(buf)
        ws = wb.active
        assert ws["B15"].fill.start_color.rgb.endswith(YELLOW)  # PN-001 阈值黄
        assert ws["B16"].fill.start_color.rgb.endswith(YELLOW)  # PN-002 回退黄
        assert ws["B17"].fill.patternType is None             # PN-003 正常不标
        assert ws["B18"].fill.start_color.rgb.endswith(RED)    # PN-004 缺货红
        wb.close()

    def test_out_of_db_fallback_rule(self, tmp_path: Path):
        """库外件号回退 库存<使用量+2，与改造前一致。"""
        src = tmp_path / "demand.xlsx"
        _build_threshold_demand(src)
        rows = read_demand(src)
        pn_cells, pn_qty = {}, {}
        for r in rows:
            pn_cells.setdefault(r.part_number, []).append(r.stock_cell)
            pn_qty.setdefault(r.part_number, []).append(r.qty)
        results = {"PN-002": 6.0, "PN-003": 5.0}  # 仅库外件号有库存
        buf, _ = write_inventory_copy(src, results, pn_qty, pn_cells, timestamp_suffix="F")
        wb = openpyxl.load_workbook(buf)
        ws = wb.active
        assert ws["B16"].fill.start_color.rgb.endswith(YELLOW)  # PN-002 5<=6<7 黄
        assert ws["B17"].fill.patternType is None             # PN-003 1<=5<3 否
        wb.close()


# ============================================================
# 集成：run_query 额外查询预警件号并回写缓存库存
# ============================================================

class TestRunQueryWarningStock:
    def test_updates_warning_stock_and_counts_demand_only(self, tmp_path: Path, monkeypatch):
        import reqman.services.connectors.amro as amro_mod
        import reqman.services.inventory_service as inv_mod

        # 需求单：PN-001 用量 2
        src = tmp_path / "demand.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A14"] = "定检专业\n（航材）"
        ws["A15"] = "发动机"; ws["B15"] = "螺钉"; ws["C15"] = "PN-001"; ws["E15"] = "2"
        wb.save(src)

        # 隔离 JsonStore：预警库含 KM-EXTRA（阈值10）
        store = JsonStore(str(tmp_path / "warn_db.json"))
        store.save_inventory_warning({"part_number": "KM-EXTRA", "name": "垫圈", "threshold": 10.0})

        async def fake_query(client, cookies, pn):
            from reqman.services.connectors.amro import KunmingStock
            return KunmingStock(5.0 if pn == "KM-EXTRA" else 1.0, "EA", "x")
        monkeypatch.setattr(amro_mod, "query_kunming_stock", fake_query)
        out_dir = tmp_path / "out"
        monkeypatch.setattr(inv_mod, "OUTPUT_DIR", out_dir)

        session = LoginSessionStore(tmp_path / "session.json", ttl_seconds=7200)
        session.save([{"name": "JSESSIONID", "value": "abc"}])
        svc = InventoryService(session, max_concurrent=2)

        dest, _, result = svc.run_query(
            src, output_stem="X",
            warning_thresholds={"KM-EXTRA": 10.0},
            warning_pns=["KM-EXTRA"], store=store,
        )
        # total 仅计需求件号
        assert result.total == 1
        assert result.shortage == 1  # PN-001 库存1 < 用量2
        # 文件落盘
        assert dest.exists()
        # 预警库件号缓存库存被回写（5.0），且需求件号不在预警库
        w = store.get_inventory_warning("KM-EXTRA")
        assert w["stock"] == 5.0
        assert store.get_inventory_warning("PN-001") is None
