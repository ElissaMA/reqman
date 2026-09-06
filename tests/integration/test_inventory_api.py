"""库存查询接口集成测试（mock AMRO）"""
import io
import re
import zipfile
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pytest


@pytest.fixture(autouse=True)
def isolated_inventory(app, tmp_path, monkeypatch):
    """隔离登录凭证缓存与输出目录到临时目录，避免污染 data/cookie/ 与 output/。"""
    from reqman.services.connectors.session import LoginSessionStore
    from reqman.services.inventory_service import InventoryService

    app.extensions["inventory_service"] = InventoryService(
        LoginSessionStore(tmp_path / "session.json", ttl_seconds=7200)
    )
    out_dir = tmp_path / "out"
    out_dir.mkdir(exist_ok=True)
    import reqman.blueprints.inventory_bp as bp_mod
    monkeypatch.setattr(bp_mod, "OUTPUT_DIR", out_dir)
    import reqman.services.inventory_service as inv_mod
    monkeypatch.setattr(inv_mod, "OUTPUT_DIR", out_dir)
    yield out_dir


def _build_demand(tmp_path: Path) -> Path:
    src = tmp_path / "demand.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A14"] = "定检专业\n（航材）"
    ws["A15"] = "发动机"; ws["B15"] = "螺钉"; ws["C15"] = "PN-001"; ws["E15"] = "2"
    wb.save(src)
    return src


def _mock_amro_query(monkeypatch, app):
    """预置登录并 mock 探活 + AMRO 查询（不发起真实网络）。"""
    svc = app.extensions["inventory_service"]
    svc.session_store.save([{"name": "JSESSIONID", "value": "abc"}])
    monkeypatch.setattr(svc, "check_login_state", lambda: "valid")
    monkeypatch.setattr(svc, "check_login", lambda: True)

    async def fake_query(client, cookies, pn):
        from reqman.services.connectors.amro import KunmingStock
        return KunmingStock(1.0, "EA", "螺钉")

    import reqman.services.connectors.amro as amro_mod
    monkeypatch.setattr(amro_mod, "query_kunming_stock", fake_query)


def _poll_inventory_status(client, timeout: float = 8) -> dict:
    """轮询 /inventory/query-status 直到 done/error（复用后台查询模式）。"""
    import time as _t

    deadline = _t.time() + timeout
    last: dict = {}
    while _t.time() < deadline:
        resp = client.get("/inventory/query-status",
                          headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json() or {}
        st = data.get("data") or {}
        last = st
        if st.get("status") in ("done", "error"):
            return st
        _t.sleep(0.05)
    return last


class TestInventoryPage:
    def test_page_accessible(self, client):
        resp = client.get("/inventory")
        assert resp.status_code == 200
        html = resp.get_data(as_text=True)
        # 「新建配置」已迁表头（base.html 链接），库存页仍可经表头触达
        assert 'href="/inventory/setup-package"' in html
        assert 'id="amroStatus"' in html
        # 查询区保留
        assert 'id="zoneDemand"' in html


class TestSession:
    def test_not_ready(self, client):
        resp = client.get("/inventory/session", headers={"X-Requested-With": "XMLHttpRequest"})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["success"] is True
        assert data["data"]["ready"] is False

    def test_cache_probe_ok_ready(self, app, client, monkeypatch):
        """缓存存在 + 探活成功 → 返回账号与登录时长。"""
        svc = app.extensions["inventory_service"]
        svc.session_store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        monkeypatch.setattr(svc, "check_login", lambda: True)
        resp = client.get("/inventory/session", headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json()
        assert data["data"]["ready"] is True
        assert data["data"]["state"] == "valid"
        assert data["data"]["account"] == "021219"
        assert "登录账号：021219已登录" in data["data"]["message"]
        assert "login_duration_seconds" in data["data"]

    def test_cache_probe_fail_expired(self, app, client, monkeypatch):
        """缓存存在 + 探活失败 → P7 登录已失效。"""
        svc = app.extensions["inventory_service"]
        svc.session_store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        monkeypatch.setattr(svc, "check_login", lambda: False)
        svc._last_probe_state = "expired"
        resp = client.get("/inventory/session", headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json()
        assert data["data"]["ready"] is False
        assert data["data"]["state"] == "expired"
        assert "重新运行登录脚本" in data["data"]["message"]

    def test_probe_error_state(self, app, client, monkeypatch):
        """网络异常 → probe_error，不清缓存、不误报未登录。"""
        svc = app.extensions["inventory_service"]
        svc.session_store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        monkeypatch.setattr(svc, "check_login", lambda: False)
        svc._last_probe_state = "probe_error"
        resp = client.get("/inventory/session", headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json()
        assert data["data"]["ready"] is False
        assert data["data"]["state"] == "probe_error"
        assert "暂不可用" in data["data"]["message"]
        assert svc.session_store.load() is not None  # 缓存保留


class TestSetupPackage:
    def test_downloads_zip(self, client):
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        assert resp.mimetype in ("application/zip", "application/x-zip-compressed")
        content = resp.data
        zf = zipfile.ZipFile(io.BytesIO(content))
        names = zf.namelist()
        assert any("amro_login.py" in n for n in names)
        assert any("start_login.bat" in n for n in names)
        assert any("README" in n for n in names)

    def _bat_source(self, resp) -> str:
        zf = zipfile.ZipFile(io.BytesIO(resp.data))
        return zf.read("start_login.bat").decode("utf-8")

    def test_bat_foreground_auto_exit(self, client):
        """ZIP 内 start_login.bat 前台运行、成功自动退出、失败传播退出码，不再要求回车。"""
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        bat = self._bat_source(resp)
        assert "start /min" not in bat
        assert ":fail" in bat
        assert "exit /b 0" in bat                       # 成功自动结束
        assert "LOGIN_EXIT" in bat and "exit /b %LOGIN_EXIT%" in bat  # 失败传播码
        assert ":done" not in bat                       # 不再保留 :done/pause 收尾

    def test_register_protocol_bat_self_locating(self, client):
        """register_protocol.bat 协议自定位（%~dp0，解压任意位置有效）+ 注册后立即启动登录，无桌面硬编码。"""
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        zf = zipfile.ZipFile(io.BytesIO(resp.data))
        bat = zf.read("register_protocol.bat").decode("utf-8")
        assert 'reg add "HKCU\\Software\\Classes\\ReqManLogin\\shell\\open\\command" /ve /d' in bat
        assert '"%~dp0start_login.bat"' in bat          # 协议指向自身目录
        assert '%%1' in bat                             # 保留 URL 参数占位
        assert "%DESK%" not in bat                      # 不再探测桌面路径
        assert "amro_login_setup" not in bat
        assert 'call "%~dp0start_login.bat"' in bat     # 注册完成后立即启动登录
        assert "exit /b %LOGIN_EXIT%" in bat            # 传播登录退出码，不再无条件 pause
        readme = zf.read("README.txt").decode("utf-8")
        assert "解压到【任意位置】" in readme
        assert "无需点击页面按钮" in readme or "无需在窗口内回车" in readme

    def test_bat_install_compat(self, client):
        """ZIP 内 start_login.bat 安装兼容性：引号 cd、新镜像、代理豁免、免 Python 装 uv、venv 实跑校验。"""
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        bat = self._bat_source(resp)
        assert 'cd /d "%~dp0"' in bat
        assert "python-build-standalone" in bat
        assert "set NO_PROXY=*" in bat
        assert "set HTTP_PROXY=" in bat
        assert "set HTTPS_PROXY=" in bat
        assert "set ALL_PROXY=" in bat
        assert ".runtime\\.installed" not in bat
        assert ":nopython" not in bat
        assert "python --version >nul 2>&1 || (py --version >nul 2>&1 || goto :nopython)" not in bat
        assert "uv.agentsmirror.com" in bat
        assert "github.com/astral-sh/uv/releases/download" in bat
        assert bat.index("uv.agentsmirror.com") < bat.index("github.com/astral-sh/uv/releases/download")
        assert "%USERPROFILE%\\.local\\bin\\uv.exe" in bat
        assert 'Get-ChildItem -Path $tmp -Recurse -Filter uv.exe' in bat
        assert '"%UV%" --version >nul 2>&1' in bat
        assert "PLAYWRIGHT_SKIP_BROWSER_DOWNLOAD=1" in bat
        assert "playwright install chromium" not in bat
        assert ".runtime\\venv\\Scripts\\python --version >nul 2>&1" in bat
        assert "python install 3.11" in bat
        assert "venv .runtime\\venv --seed" in bat
        assert "venv .runtime\\venv\r\n" in bat
        assert "pypi.tuna.tsinghua.edu.cn" in bat
        assert 'python -c "import playwright"' in bat
        assert "--force-reinstall playwright" in bat

    def _py_source(self, resp) -> str:
        zf = zipfile.ZipFile(io.BytesIO(resp.data))
        return zf.read("amro_login.py").decode("utf-8")

    def test_injects_fallback_host_url(self, client, monkeypatch):
        """未配置 AMRO_PUBLIC_URL → 回退当前访问地址（host_url）。"""
        import reqman.blueprints.inventory_bp as bp_mod
        monkeypatch.setattr(bp_mod, "AMRO_PUBLIC_URL", "")
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        src = self._py_source(resp)
        assert 'SERVER_URL = "http://localhost"' in src
        assert 'UPLOAD_URL = "http://localhost/inventory/login/upload"' in src
        assert "trust_env=False" in src

    def test_py_uses_system_browser_chrome_to_msedge(self, client):
        """amro_login.py 三级回退：Chrome → Edge → 显式 Edge 路径，弃用 chromium 下载。"""
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        src = self._py_source(resp)
        assert '{"channel": "chrome"}' in src
        assert '{"channel": "msedge"}' in src
        assert src.index('"chrome"') < src.index('"msedge"')
        assert "executable_path" in src
        assert "ProgramFiles(x86)" in src
        assert "msedge.exe" in src
        assert "未检测到 Chrome/Edge，请安装浏览器后重试" in src

    def test_py_login_version_and_auto_probe(self, client):
        """amro_login.py 注入 LOGIN_VERSION=4，且为自动探活上传（无回车/无完成按钮）。"""
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        src = self._py_source(resp)
        assert 'LOGIN_VERSION = "4"' in src
        assert "input(" not in src                      # 不再窗口回车
        assert "_wait_confirm" not in src
        assert "__reqman_done_btn" not in src           # 不再注入完成按钮
        assert "MM_PARTNUMBERCHAXUN_LIST" in src        # 自动只读探活
        assert "ACCOUNT_SELECTORS" in src               # 从页面读取账号
        assert "raise SystemExit" in src                # 失败传播非零退出码

    def test_injects_configured_public_url(self, client, monkeypatch):
        """配置 AMRO_PUBLIC_URL → 注入配置的公网地址。"""
        import reqman.blueprints.inventory_bp as bp_mod
        monkeypatch.setattr(bp_mod, "AMRO_PUBLIC_URL", "http://8.137.15.167")
        resp = client.get("/inventory/setup-package")
        assert resp.status_code == 200
        src = self._py_source(resp)
        assert 'SERVER_URL = "http://8.137.15.167"' in src
        assert 'UPLOAD_URL = "http://8.137.15.167/inventory/login/upload"' in src
        assert "trust_env=False" in src


class TestCheckConfig:
    def test_rejects_wrong_content(self, client):
        resp = client.post(
            "/inventory/check-config",
            data={"file": (io.BytesIO(b"not a script"), "foo.txt")},
            headers={"X-Requested-With": "XMLHttpRequest"},
            content_type="multipart/form-data",
        )
        assert resp.status_code == 400
        data = resp.get_json()
        assert data["success"] is False


class TestLoginUpload:
    def test_upload_saves_with_account(self, app, client):
        cookies_json = '[{"name":"JSESSIONID","value":"abc123"}]'
        resp = client.post(
            "/inventory/login/upload",
            data={"cookies": cookies_json, "account": "021219"},
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 200
        assert resp.get_json()["success"] is True
        loaded = app.extensions["inventory_service"].session_store.load()
        assert loaded["account"] == "021219"

    def test_upload_requires_account(self, client):
        resp = client.post(
            "/inventory/login/upload",
            data={"cookies": '[{"name":"JSESSIONID","value":"abc"}]'},
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 400

    def test_upload_requires_jsessionid(self, client):
        resp = client.post(
            "/inventory/login/upload",
            data={"cookies": '[{"name":"OTHER","value":"x"}]', "account": "021219"},
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 400

    def test_upload_invalid_json(self, client):
        resp = client.post(
            "/inventory/login/upload",
            data={"cookies": "not json", "account": "021219"},
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 400


class TestQuery:
    def test_query_requires_login(self, client, tmp_path):
        src = _build_demand(tmp_path)
        with src.open("rb") as f:
            resp = client.post(
                "/inventory/query",
                data={"file": (f, "demand.xlsx")},
                headers={"X-Requested-With": "XMLHttpRequest"},
                content_type="multipart/form-data",
            )
        assert resp.status_code == 401
        data = resp.get_json()
        assert data["success"] is False
        assert "重新运行登录脚本" in data["message"]


class TestDownloadFilenameEncoding:
    """下载链接文件名 URL 编码：含 + 空格 & 的文件名必须经 quote 命中磁盘真实文件。"""

    def test_download_with_special_chars_via_quoted_link(self, client, tmp_path, monkeypatch):
        from urllib.parse import quote

        import reqman.blueprints.inventory_bp as bp_mod
        out_dir = tmp_path / "out_dl"
        out_dir.mkdir(parents=True, exist_ok=True)
        monkeypatch.setattr(bp_mod, "OUTPUT_DIR", out_dir)
        name = "定检需求单（B-6836 68A+24MO）2026.09.11_库存已填_20260904_152506.xlsx"
        (out_dir / name).write_bytes(b"fake-xlsx")
        resp = client.get(f"/inventory/download?file={quote(name)}")
        assert resp.status_code == 200
        assert resp.data == b"fake-xlsx"

    def test_download_unencoded_plus_misses_file(self, client, tmp_path, monkeypatch):
        """未编码的 + 被 WSGI 解码为空格 → 文件名不符 → 拒绝（证明编码必要）。"""
        import reqman.blueprints.inventory_bp as bp_mod
        out_dir = tmp_path / "out_dl2"
        out_dir.mkdir(parents=True, exist_ok=True)
        monkeypatch.setattr(bp_mod, "OUTPUT_DIR", out_dir)
        real = "68A+24MO_库存已填.xlsx"
        (out_dir / real).write_bytes(b"fake")
        resp = client.get(
            "/inventory/download?file=68A+24MO_库存已填.xlsx",
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code in (400, 404)  # 命中的是 "68A 24MO_库存已填.xlsx"（不存在）

    def test_reconcile_keeps_encoded_link_when_file_exists(self, monkeypatch):
        import tempfile
        from urllib.parse import quote

        import reqman.blueprints.inventory_bp as bp_mod
        with tempfile.TemporaryDirectory() as td:
            out_dir = Path(td)
            monkeypatch.setattr(bp_mod, "OUTPUT_DIR", out_dir)
            name = "68A+24MO_库存已填_20260904.xlsx"
            (out_dir / name).write_bytes(b"x")
            last = {"download_url": f"/inventory/download?file={quote(name)}"}
            out = bp_mod._reconcile_inventory_download(dict(last))
            assert out["download_url"] == last["download_url"]  # 保留原链接

    def test_reconcile_falls_back_to_latest_and_quotes(self, monkeypatch):
        import tempfile
        from urllib.parse import quote

        import reqman.blueprints.inventory_bp as bp_mod
        with tempfile.TemporaryDirectory() as td:
            out_dir = Path(td)
            monkeypatch.setattr(bp_mod, "OUTPUT_DIR", out_dir)
            latest = "定检需求单（B-6836 68A+24MO）_库存已填_20260905.xlsx"
            (out_dir / latest).write_bytes(b"x")
            stale = {"download_url": "/inventory/download?file=old_库存已填.xlsx"}
            out = bp_mod._reconcile_inventory_download(dict(stale))
            assert out["download_url"] == f"/inventory/download?file={quote(latest)}"

    def test_reconcile_no_files_no_link(self, monkeypatch):
        import tempfile

        import reqman.blueprints.inventory_bp as bp_mod
        with tempfile.TemporaryDirectory() as td:
            monkeypatch.setattr(bp_mod, "OUTPUT_DIR", Path(td))
            out = bp_mod._reconcile_inventory_download({"download_url": "/inventory/download?file=gone.xlsx"})
            assert out["download_url"] == ""


class TestRequireAmroSessionStates:
    """前置探活三态：probe_error 放行（避免网络抖动误判未登录），none/expired 阻断。"""

    def test_valid_allows_query(self, app, monkeypatch):
        from reqman.services import amro_sync

        svc = app.extensions["inventory_service"]
        monkeypatch.setattr(svc, "check_login_state", lambda: "valid")
        with app.test_request_context():
            assert amro_sync.require_amro_session() is True

    def test_probe_error_allows_query(self, app, monkeypatch):
        from reqman.services import amro_sync

        svc = app.extensions["inventory_service"]
        monkeypatch.setattr(svc, "check_login_state", lambda: "probe_error")
        with app.test_request_context():
            assert amro_sync.require_amro_session() is True

    def test_expired_blocks_query(self, app, monkeypatch):
        from reqman.services import amro_sync

        svc = app.extensions["inventory_service"]
        monkeypatch.setattr(svc, "check_login_state", lambda: "expired")
        with app.test_request_context():
            assert amro_sync.require_amro_session() is False

    def test_none_blocks_query(self, app, monkeypatch):
        from reqman.services import amro_sync

        svc = app.extensions["inventory_service"]
        monkeypatch.setattr(svc, "check_login_state", lambda: "none")
        with app.test_request_context():
            assert amro_sync.require_amro_session() is False


class TestQueryFlow:
    def test_query_success_returns_stats_and_file(self, app, client, tmp_path, monkeypatch, isolated_inventory):
        """成功路径：启动后台任务 → 轮询至 done → 统计 + 文件名 + 输出落盘 output/。"""
        _mock_amro_query(monkeypatch, app)
        out_dir = isolated_inventory

        src = _build_demand(tmp_path)
        with src.open("rb") as f:
            resp = client.post(
                "/inventory/query",
                data={"file": (f, "demand.xlsx")},
                headers={"X-Requested-With": "XMLHttpRequest"},
                content_type="multipart/form-data",
            )
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["success"] is True
        assert data["data"]["started"] is True

        status = _poll_inventory_status(client)
        assert status.get("status") == "done", status
        s = status["summary"]
        assert s["total"] == 1
        assert s["success"] == 1
        assert s["shortage"] == 1
        assert s["filename"].startswith("demand_库存已填_")
        # 文件已写入 output/
        saved = list(out_dir.glob("demand_库存已填_*.xlsx"))
        assert len(saved) == 1
        wb = openpyxl.load_workbook(saved[0])
        assert wb.active["G15"].value == 1
        wb.close()

    def test_query_busy_conflict(self, app, client, tmp_path, monkeypatch, isolated_inventory):
        """全局互斥：已有查询在跑 → 409 + busy 文案（不排队）。"""
        import time as _time

        from reqman.services import amro_sync
        _mock_amro_query(monkeypatch, app)
        deadline = _time.time() + 3
        while not amro_sync.try_begin_query("查询飞机数据"):
            if _time.time() > deadline:
                pytest.fail("query slot still busy")
            _time.sleep(0.02)
        try:
            src = _build_demand(tmp_path)
            with src.open("rb") as f:
                resp = client.post(
                    "/inventory/query",
                    data={"file": (f, "demand.xlsx")},
                    headers={"X-Requested-With": "XMLHttpRequest"},
                    content_type="multipart/form-data",
                )
            assert resp.status_code == 409
            msg = resp.get_json()["message"]
            assert "已有查询任务进行中" in msg and "查询飞机数据" in msg
        finally:
            amro_sync.end_query()

    def test_query_clears_old_staging(self, app, client, tmp_path, monkeypatch, isolated_inventory):
        """上传新需求单查询 → 旧的 *_库存已填_* 暂存被清除，output/ 仅存最新（后台完成后）。"""
        _mock_amro_query(monkeypatch, app)
        out_dir = isolated_inventory
        (out_dir / "旧需求单_库存已填_20260101_000000.xlsx").write_bytes(b"old")
        (out_dir / "无关文件.txt").write_bytes(b"keep")

        src = _build_demand(tmp_path)
        with src.open("rb") as f:
            resp = client.post(
                "/inventory/query",
                data={"file": (f, "demand.xlsx")},
                headers={"X-Requested-With": "XMLHttpRequest"},
                content_type="multipart/form-data",
            )
        assert resp.status_code == 200
        status = _poll_inventory_status(client)
        assert status.get("status") == "done", status
        staged = [p.name for p in out_dir.glob("*_库存已填_*.xlsx")]
        assert len(staged) == 1  # 旧暂存已清，仅新产物
        assert "旧需求单_库存已填_" not in staged[0]
        assert (out_dir / "无关文件.txt").exists()  # 未误删其他文件
        # 上传暂存 .staging_* 已在 job 内清理，不残留
        assert not list(out_dir.glob(".staging_*.xlsx"))


class TestDownload:
    def test_download_ok(self, client, tmp_path, isolated_inventory):
        out_dir = isolated_inventory
        f = out_dir / "需求单A_库存已填_20260101_000000.xlsx"
        f.write_bytes(b"xlsx-bytes")
        resp = client.get("/inventory/download?file=" + f.name)
        assert resp.status_code == 200
        assert resp.data == b"xlsx-bytes"

    def test_download_rejects_path_traversal(self, client, isolated_inventory):
        resp = client.get(
            "/inventory/download?file=..%2F..%2Fsecret.txt",
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 400
        resp2 = client.get(
            "/inventory/download?file=..%5Csecret.txt",
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp2.status_code == 400

    def test_download_not_exists(self, client, isolated_inventory):
        resp = client.get(
            "/inventory/download?file=nonexist_库存已填_1.xlsx",
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 400


class TestLatestOutput:
    def test_no_output(self, client, isolated_inventory):
        resp = client.get("/inventory/latest-output", headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json()
        assert data["data"]["exists"] is False

    def test_has_latest(self, client, isolated_inventory):
        out_dir = isolated_inventory
        (out_dir / "A_库存已填_1.xlsx").write_bytes(b"a")
        (out_dir / "B_库存已填_2.xlsx").write_bytes(b"b")
        resp = client.get("/inventory/latest-output", headers={"X-Requested-With": "XMLHttpRequest"})
        data = resp.get_json()
        assert data["data"]["exists"] is True
        assert data["data"]["filename"] == "B_库存已填_2.xlsx"


import reqman.services.amro_sync as amro_sync_mod


class TestInventoryLastQueryDisplay:
    def test_index_no_duplicate_resultbox(self, client):
        """页面不应再有第二套 #resultBox / loadLatestOutput / renderResult 输出。"""
        resp = client.get("/inventory")
        html = resp.get_data(as_text=True)
        assert 'id="resultBox"' not in html
        assert "loadLatestOutput" not in html
        assert "renderResult" not in html

    def test_index_renders_last_query_block(self, app, client, tmp_path, monkeypatch, isolated_inventory):
        """跑一次查询后，库存页仅有一处「上次查询」摘要+可用下载链接。"""
        out_dir = isolated_inventory
        monkeypatch.setattr(amro_sync_mod, "OUTPUT_DIR", out_dir)
        _mock_amro_query(monkeypatch, app)
        src = _build_demand(tmp_path)
        with src.open("rb") as f:
            resp = client.post(
                "/inventory/query",
                data={"file": (f, "demand.xlsx")},
                headers={"X-Requested-With": "XMLHttpRequest"},
                content_type="multipart/form-data",
            )
        assert resp.get_json()["success"] is True
        status = _poll_inventory_status(client)
        assert status.get("status") == "done", status

        resp = client.get("/inventory")
        html = resp.get_data(as_text=True)
        assert "上次查询（" in html
        assert 'id="lastQueryBlock"' in html
        assert "⬇下载库存查询" in html
        m = re.search(r'href="(/inventory/download\?file=[^"]+)"', html)
        assert m, "摘要块未渲染下载链接"
        dl = client.get(m.group(1))
        assert dl.status_code == 200

    def test_stale_download_reconciled_to_live(self, client, monkeypatch, isolated_inventory):
        """持久化摘要指向已删除的静态文件时，下载链接应回退到当前实时最新文件。"""
        out_dir = isolated_inventory
        monkeypatch.setattr(amro_sync_mod, "OUTPUT_DIR", out_dir)
        real = out_dir / "需求单_库存已填_20260101_000000.xlsx"
        real.write_bytes(b"xlsx-bytes")
        (out_dir / "last_query_inventory_query.json").write_text(
            '{"label":"查询库存","finished_at":"2026-01-01 00:00:00",'
            '"summary":"查询完成：共 1 件号，成功 1，失败 0，标红 1，标黄 0",'
            '"download_url":"/inventory/download?file=已删除_库存已填_20250101_000000.xlsx"}',
            encoding="utf-8",
        )
        resp = client.get("/inventory")
        html = resp.get_data(as_text=True)
        assert "上次查询（" in html
        m = re.search(r'href="(/inventory/download\?file=[^"]+)"', html)
        assert m, "摘要块未渲染下载链接"
        assert "已删除_库存已填" not in unquote(m.group(1))
        assert "需求单_库存已填" in unquote(m.group(1))  # 回退链接的文件名已 URL 编码
        dl = client.get(m.group(1))  # 编码链接可直接下载命中真实文件
        assert dl.status_code == 200
        assert dl.data == b"xlsx-bytes"
