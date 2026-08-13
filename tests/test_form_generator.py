"""form_generator 3.2.4 改动测试：数量空输出、category 兜底移除（不KeyError）"""
import openpyxl
import pytest

from reqman.services.form_generator import _write_category_block_min_rows


class TestWriteCategoryBlockMinRows:
    """_write_category_block_min_rows：3.2.4 改动回归"""

    def _run(self, items, **kw):
        wb = openpyxl.Workbook()
        ws = wb.active
        next_row = _write_category_block_min_rows(
            ws, 1, items, min_rows=1,
            name_key="device_name", task_key="task_name", **kw,
        )
        return ws, next_row

    def test_quantity_empty_output(self):
        """3.2.4：数量为空 → Excel 数量列输出空（不默认"1"）"""
        ws, _ = self._run([{
            "category": "发动机",
            "device_name": "扳手",
            "part_number": "PN-1",
            "quantity": "",
            "remark": "",
        }])
        assert ws.cell(row=1, column=5).value in (None, "")

    def test_quantity_kept_when_present(self):
        """数量有值 → 原样输出"""
        ws, _ = self._run([{
            "category": "发动机",
            "device_name": "扳手",
            "part_number": "PN-1",
            "quantity": "3",
            "remark": "",
        }])
        assert ws.cell(row=1, column=5).value == "3"

    def test_category_missing_raises_keyerror(self):
        """3.2.4 移除 category 兜底（原默认"电子"）：item 缺 category → KeyError

        数据源（工卡/工卡组）现在保证 category 必填，移除兜底后正常路径不 KeyError。
        """
        with pytest.raises(KeyError):
            self._run([{"device_name": "扳手"}])

    def test_normal_flow_no_keyerror(self):
        """正常数据（全部含 category）→ 不 KeyError，按 发动机/机体/电子 三组渲染"""
        ws, next_row = self._run([
            {"category": "发动机", "device_name": "扳手", "part_number": "", "quantity": "", "remark": ""},
            {"category": "电子", "device_name": "万用表", "part_number": "", "quantity": "2", "remark": ""},
        ])
        assert next_row > 1
        # 行1=发动机（扳手），行2=机体（暂无），行3=电子（万用表）
        assert ws.cell(row=1, column=1).value == "发动机"
        assert ws.cell(row=1, column=2).value == "扳手"
        assert ws.cell(row=2, column=1).value == "机体"
        assert ws.cell(row=3, column=1).value == "电子"
        assert ws.cell(row=3, column=5).value == "2"


class TestExcelRemarkDisplay:
    """备注显示统一：Excel 中"检查有问题领用"用中文逗号拼接备注"""

    def _run(self, items, **kw):
        wb = openpyxl.Workbook()
        ws = wb.active
        _write_category_block_min_rows(
            ws, 1, items, min_rows=1,
            name_key="device_name", task_key="task_name", **kw,
        )
        return ws

    def test_check_issue_with_remark_chinese_comma(self):
        """有问题领用+备注 → '检查有问题领用，备注'（中文逗号分隔）"""
        ws = self._run([{
            "category": "发动机", "device_name": "扳手", "part_number": "",
            "quantity": "", "remark": "已磨损",
            "usage_type": "检查有问题领用",
        }])
        assert ws.cell(row=1, column=6).value == "检查有问题领用，已磨损"

    def test_check_issue_empty_remark_only_type(self):
        """有问题领用+空备注 → 只显示'检查有问题领用'"""
        ws = self._run([{
            "category": "发动机", "device_name": "扳手", "part_number": "",
            "quantity": "", "remark": "",
            "usage_type": "检查有问题领用",
        }])
        assert ws.cell(row=1, column=6).value == "检查有问题领用"

    def test_must_use_only_remark(self):
        """必须使用 → 只显示备注（不加类型前缀）"""
        ws = self._run([{
            "category": "发动机", "device_name": "扳手", "part_number": "",
            "quantity": "", "remark": "随机携带",
            "usage_type": "必须使用",
        }])
        assert ws.cell(row=1, column=6).value == "随机携带"

    def test_no_usage_type_only_remark(self):
        """无 usage_type → 只显示备注"""
        ws = self._run([{
            "category": "发动机", "device_name": "扳手", "part_number": "",
            "quantity": "", "remark": "备用",
        }])
        assert ws.cell(row=1, column=6).value == "备用"

    def test_no_usage_type_empty_remark_blank(self):
        """无 usage_type 且无备注 → 空"""
        ws = self._run([{
            "category": "发动机", "device_name": "扳手", "part_number": "",
            "quantity": "", "remark": "",
        }])
        assert ws.cell(row=1, column=6).value in (None, "")
