"""库存查询业务编排测试"""
from pathlib import Path

import pytest

from reqman.services.connectors.session import LoginSessionStore
from reqman.services.inventory_service import InventoryService


def _store(tmp_path: Path):
    return LoginSessionStore(tmp_path / "session.json", ttl_seconds=7200)


class TestGetLoginStatus:
    def test_not_ready(self, tmp_path: Path):
        status = InventoryService(_store(tmp_path)).get_login_status()
        assert status["ready"] is False
        assert status["state"] == "none"
        assert status["account"] is None
        assert status["login_duration_seconds"] == 0

    def test_ready_with_probe(self, tmp_path: Path, monkeypatch):
        store = _store(tmp_path)
        store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        svc = InventoryService(store)
        monkeypatch.setattr(svc, "check_login", lambda: True)
        status = svc.get_login_status()
        assert status["ready"] is True
        assert status["state"] == "valid"
        assert status["account"] == "021219"
        assert status["login_duration_seconds"] >= 0

    def test_probe_failed_not_ready(self, tmp_path: Path, monkeypatch):
        store = _store(tmp_path)
        store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        svc = InventoryService(store)
        monkeypatch.setattr(svc, "check_login", lambda: False)
        svc._last_probe_state = "probe_error"
        status = svc.get_login_status()
        assert status["ready"] is False
        assert status["state"] == "probe_error"


class TestProbeCache:
    """探活短 TTL 缓存：TTL 内重复探活复用结论，不再真实外呼耗线程"""

    def _svc_with_cookie(self, tmp_path: Path):
        store = _store(tmp_path)
        store.save([{"name": "JSESSIONID", "value": "abc"}], account="021219")
        return InventoryService(store)

    def test_ttl_window_reuses_probe(self, tmp_path: Path, monkeypatch):
        svc = self._svc_with_cookie(tmp_path)
        calls = []

        async def fake_probe(data):
            calls.append(1)
            return "valid"

        monkeypatch.setattr(svc, "_probe_cached_session", fake_probe)
        assert svc.check_login_state() == "valid"
        assert svc.check_login_state() == "valid"   # TTL 内复用
        assert svc.check_login_state() == "valid"
        assert len(calls) == 1                       # 只真实探活一次

    def test_new_login_invalidates_cache(self, tmp_path: Path, monkeypatch):
        svc = self._svc_with_cookie(tmp_path)
        calls = []

        async def fake_probe(data):
            calls.append(1)
            return "valid"

        monkeypatch.setattr(svc, "_probe_cached_session", fake_probe)
        svc.check_login_state()
        # 重新登录（login_at 变化）→ 缓存失效，需重新探活
        svc.session_store.save([{"name": "JSESSIONID", "value": "new"}], account="021219")
        assert svc.check_login_state() == "valid"
        assert len(calls) == 2

    def test_no_cookie_skips_probe(self, tmp_path: Path, monkeypatch):
        svc = InventoryService(_store(tmp_path))
        async def fail_probe(data):  # pragma: no cover - 不应被调用
            raise AssertionError("无凭证不应探活")
        monkeypatch.setattr(svc, "_probe_cached_session", fail_probe)
        assert svc.check_login_state() == "none"


class TestSaveLogin:
    def test_saves_and_ready(self, tmp_path: Path):
        store = _store(tmp_path)
        svc = InventoryService(store)
        svc.save_login([{"name": "JSESSIONID", "value": "xyz"}], account="021219")
        loaded = store.load()
        assert loaded["cookies"]["JSESSIONID"] == "xyz"
        assert loaded["account"] == "021219"


class TestRunQuery:
    def test_query_flow(self, tmp_path: Path, monkeypatch):
        import openpyxl
        src = tmp_path / "demand.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A14"] = "定检专业\n（航材）"
        ws["A15"] = "发动机"; ws["B15"] = "螺钉"; ws["C15"] = "PN-001"; ws["E15"] = "2"
        wb.save(src)

        async def fake_query(client, cookies, pn):
            from reqman.services.connectors.amro import KunmingStock
            return KunmingStock(1.0, "EA", "螺钉")

        import reqman.services.connectors.amro as amro_mod
        monkeypatch.setattr(amro_mod, "query_kunming_stock", fake_query)

        out_dir = tmp_path / "out"
        import reqman.services.inventory_service as inv_mod
        monkeypatch.setattr(inv_mod, "OUTPUT_DIR", out_dir)

        store = _store(tmp_path)
        store.save([{"name": "JSESSIONID", "value": "abc"}])
        svc = InventoryService(store, max_concurrent=2)
        dest, filename, result = svc.run_query(src, output_stem="原始需求单")
        assert dest.exists()
        assert result.total == 1
        assert result.success == 1
        assert result.shortage == 1  # 需求2 库存1 → 标红
        assert result.filename == filename
        # 文件名保留原名 stem
        assert filename.startswith("原始需求单_库存已填_")
        # 副本可被 openpyxl 打开
        ws2 = openpyxl.load_workbook(dest).active
        assert ws2["G15"].value == 1

    def test_session_expiry_aborts_batch_without_output_or_warning_write(self, tmp_path, monkeypatch):
        import openpyxl

        src = tmp_path / "expiry.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A14"] = "定检专业\n（航材）"
        ws["A15"] = "发动机"; ws["B15"] = "螺钉"; ws["C15"] = "PN-EXPIRE"; ws["E15"] = "1"
        wb.save(src)

        import reqman.services.connectors.amro as amro_mod
        import reqman.services.inventory_service as inv_mod
        from reqman.services.connectors.amro import AmroSessionExpired

        async def expired_query(client, cookies, pn):
            raise AmroSessionExpired("session expired")

        monkeypatch.setattr(amro_mod, "query_kunming_stock", expired_query)
        out_dir = tmp_path / "out"
        monkeypatch.setattr(inv_mod, "OUTPUT_DIR", out_dir)
        warning_store = _WarningStoreSpy()
        store = _store(tmp_path)
        store.save([{"name": "JSESSIONID", "value": "abc"}])
        svc = InventoryService(store, max_concurrent=2)

        with pytest.raises(AmroSessionExpired):
            svc.run_query(
                src,
                output_stem="expiry",
                warning_thresholds={"PN-EXPIRE": 2},
                warning_pns=["PN-EXPIRE"],
                store=warning_store,
            )
        assert not out_dir.exists() or list(out_dir.iterdir()) == []
        assert warning_store.saved == []


class _WarningStoreSpy:
    def __init__(self):
        self.saved = []

    def get_inventory_warning(self, part_number):
        return {"part_number": part_number}

    def save_inventory_warning(self, data):
        self.saved.append(data)
