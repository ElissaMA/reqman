"""借用清单生成器测试（统一重复判定：件号→名称 两级；分区/通用取最大/排除预印与常用外场/扩行）"""
import openpyxl

from reqman.services.checklist_generator import (
    _group_items,
    _merge_group,
    _same_item,
    generate_chemical_list,
    generate_tool_list,
)


def _tool(name, category, qty="1", pn="PN-1", usage_type="必须使用"):
    return {"device_name": name, "part_number": pn, "quantity": qty,
            "remark": "", "usage_type": usage_type, "category": category}


def _mat(name, qty="1", pn="PN-M", remark="优先使用开封航化", usage_type="必须使用"):
    return {"material_name": name, "part_number": pn, "quantity": qty,
            "remark": remark, "usage_type": usage_type, "category": "电子"}


def _load_tool(buffer):
    return openpyxl.load_workbook(buffer)["定检工具"]


def _load_chem(buffer):
    return openpyxl.load_workbook(buffer)["定检航化"]


class TestDedupStandard:
    def test_same_item(self):
        a = {"part_number": "PN-1", "device_name": "扳手"}
        assert _same_item(a, {"part_number": "pn-1", "device_name": "X"}, "device_name")   # 件号相同
        assert _same_item(a, {"part_number": "PN-2", "device_name": "扳手"}, "device_name")  # 名称相同
        assert not _same_item(a, {"part_number": "PN-2", "device_name": "螺丝刀"}, "device_name")
        # 件号为空 → 直接比名称
        empty = {"part_number": "", "device_name": "扳手"}
        assert _same_item(empty, {"part_number": "PN-9", "device_name": "扳手"}, "device_name")
        assert not _same_item(empty, {"part_number": "PN-9", "device_name": "钳"}, "device_name")

    def test_group_and_merge(self):
        rows = [
            _tool("力矩扳手", "发动机", "2", pn="PN-A"),
            _tool("力矩扳手", "发动机", "5", pn="PN-B"),   # 名称相同件号不同 → 同组
            _tool("专用塞尺", "发动机", "1", pn="PN-C"),   # 专用前缀：名称不同 → 不合并
            _tool("塞尺", "发动机", "1", pn="PN-D"),
        ]
        groups = _group_items(rows, "device_name")
        assert len(groups) == 3
        assert _merge_group(groups[0])["quantity"] == "5"   # 取最大
        assert _merge_group(groups[0])["part_number"] == "PN-B"

    def test_merge_unparseable_qty_fallback(self):
        rows = [_tool("卡尺", "电子", "", pn="P1"), _tool("卡尺", "电子", "若干", pn="P2")]
        group = _group_items(rows, "device_name")
        assert len(group) == 1
        assert _merge_group(group[0])["quantity"] == "若干"


class TestGenerateToolList:
    def _form(self):
        return {"reg": "B-1234", "description": "46A", "date": "2026-09-06"}

    def test_single_major_tools_into_zone(self):
        tools = [_tool("内六角扳手", "发动机", "2", pn="PN-1"),
                 _tool("静电手环", "电子", "1", pn="PN-2")]
        buffer, filename = generate_tool_list(self._form(), tools)
        assert "定检中队零散工具借用清单（B-1234 46A）2026-09-06.xlsx" == filename
        ws = _load_tool(buffer)
        assert ws["B19"].value == "内六角扳手"    # 发动机区首行
        assert ws["D19"].value == "2"
        assert ws["B39"].value == "静电手环"      # 电子区首行

    def test_same_pn_different_name_merged_into_generic(self):
        tools = [_tool("孔探仪", "发动机", "1", pn="IAE6F10408"),
                 _tool("孔探仪带组件", "机体", "2", pn="IAE6F10408")]   # 同件号 → 同组
        buffer, _ = generate_tool_list(self._form(), tools)
        ws = _load_tool(buffer)
        assert ws["B16"].value == "孔探仪"        # 两专业 → 通用区
        assert ws["D16"].value == "2"             # 取最大
        assert ws["B19"].value is None           # 专业区不重复

    def test_same_name_different_pn_merged_into_generic(self):
        tools = [_tool("专用吊具", "发动机", "1"), _tool("专用吊具", "机体", "2")]
        buffer, _ = generate_tool_list(self._form(), tools)
        ws = _load_tool(buffer)
        assert ws["B16"].value == "专用吊具"
        assert ws["D16"].value == "2"
        assert ws["B19"].value is None

    def test_scope_and_cross_major_skipped(self):
        tools = [_tool("孔探仪", "发动机", "1", pn="IAE-1"),
                 _tool("深度检测仪", "机体", "1", pn="DNS-2"),   # 件号名称均不同 → 不合并
                 _tool("特检工具", "特检", "1", pn="PN-S"),
                 _tool("支援工具", "支援", "1", pn="PN-Z"),
                 _tool("备用扳手", "发动机", "1", pn="PN-X", usage_type="检查有问题领用")]
        buffer, _ = generate_tool_list(self._form(), tools)
        ws = _load_tool(buffer)
        written = [ws.cell(row=r, column=2).value for r in range(16, 49)]
        assert written.count("孔探仪") == 1            # 各写一行，不合并
        assert written.count("深度检测仪") == 1
        assert "特检工具" not in written and "支援工具" not in written
        assert "备用扳手" not in written

    def test_skip_when_hits_generic_preprinted(self):
        tools = [_tool("手套", "机体", "5")]                        # 名称命中预印
        buffer, _ = generate_tool_list(self._form(), tools)
        ws = _load_tool(buffer)
        assert ws["B29"].value is None            # 专业区不写
        assert ws["D3"].value == 40               # 预印行不被覆盖（模板数值单元格）

    def test_skip_when_pn_hits_generic_preprinted(self):
        # 名称命中预印「毛巾」且件号不同 → 仍按名称判重跳过
        tools = [_tool("毛巾", "机体", "3", pn="PN-OTHER")]
        buffer, _ = generate_tool_list(self._form(), tools)
        ws = _load_tool(buffer)
        assert ws["B29"].value is None or ws["B29"].value != "毛巾"

    def test_zone_full_extends_rows(self):
        tools = [_tool(f"专用吊具{i}", "发动机", "1", pn=f"PN-T{i}") for i in range(5)]
        tools += [_tool(f"专用吊具{i}", "机体", "1", pn=f"PN-T{i}") for i in range(5)]
        buffer, _ = generate_tool_list(self._form(), tools)
        wb = openpyxl.load_workbook(buffer)
        ws = wb["定检工具"]
        merges = {str(r) for r in ws.merged_cells.ranges}
        assert "A3:A20" in merges                 # 通用区扩至 3-20（+2）
        assert "A21:A30" in merges                # 发动机区整体下移 2
        for i in range(5):
            assert ws.cell(row=16 + i, column=2).value == f"专用吊具{i}"


class TestGenerateChemicalList:
    def _form(self):
        return {"reg": "B-1234", "description": "46A", "date": "2026-09-06"}

    def test_opened_remark_required(self):
        materials = [
            _mat("密封胶 MQ-1", "2"),
            _mat("普通油脂", "3", remark=""),
        ]
        buffer, filename = generate_chemical_list(self._form(), materials)
        assert "定检中队开封航化借用清单（B-1234 46A）2026-09-06.xlsx" == filename
        ws = _load_chem(buffer)
        assert ws["B18"].value == "密封胶 MQ-1"   # 非例行区首行
        assert ws["B19"].value is None            # 非开封航化不写

    def test_same_pn_different_name_merged(self):
        materials = [
            _mat("粘接剂", "1", pn="LOCTITE242"),
            _mat("胶", "3", pn="LOCTITE242"),     # 同件号 → 同组
        ]
        buffer, _ = generate_chemical_list(self._form(), materials)
        ws = _load_chem(buffer)
        assert ws["B18"].value == "粘接剂"        # 组代表=首行
        assert ws["D18"].value == "3"             # 取最大

    def test_same_name_different_pn_merged(self):
        materials = [
            _mat("除冰剂", "2", pn="PN-A"),
            _mat("除冰剂", "1", pn="PN-B"),
        ]
        buffer, _ = generate_chemical_list(self._form(), materials)
        ws = _load_chem(buffer)
        assert ws["B18"].value == "除冰剂"
        assert ws["D18"].value == "2"             # 取最大

    def test_empty_pn_falls_back_to_name(self):
        materials = [
            _mat("特种油脂", "1", pn=""),
            _mat("特种油脂", "4", pn=""),
        ]
        buffer, _ = generate_chemical_list(self._form(), materials)
        ws = _load_chem(buffer)
        assert ws["B18"].value == "特种油脂"
        assert ws["D18"].value == "4"

    def test_exclude_common_and_field_zones(self):
        materials = [
            _mat("纸胶带（宽）", "5"),                     # 名称命中常用预印
            _mat("清洗剂", "2", pn="ZOK27"),               # 件号命中外场预印
            _mat("75%酒精清洗剂", "1", pn="75JJ"),         # 非预印 → 照写
        ]
        buffer, _ = generate_chemical_list(self._form(), materials)
        ws = _load_chem(buffer)
        names = [ws.cell(row=r, column=2).value for r in range(18, 25)]
        assert names[0] == "75%酒精清洗剂"
        assert "纸胶带（宽）" not in names
        assert "清洗剂" not in names

    def test_exclude_spare_and_invalid_usage(self):
        materials = [
            _mat("备用密封胶", "2", usage_type="检查有问题领用"),
            _mat("开封航化正常件", "1"),
        ]
        buffer, _ = generate_chemical_list(self._form(), materials)
        ws = _load_chem(buffer)
        assert ws["B18"].value == "开封航化正常件"
        assert ws["B19"].value is None
