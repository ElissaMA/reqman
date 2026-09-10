"""模板契约快照测试（修复计划阶段0）。

固定区文字/合并/打印区按 assets 真实模板断言；任何生成器写入模板固定区的行为
都会在此被拦截。动态区白名单与结构异常报错行为也在此验证。
"""
import openpyxl
import pytest

from reqman.config import (
    CHECK_TEMPLATE_FILE,
    MATERIALS_TEMPLATE_FILE,
    REMINDER_TEMPLATE_FILE,
    TEMPLATE_FILE,
    TOOLS_TEMPLATE_FILE,
)


def _cell(ws, coord):
    return str(ws[coord].value or "").strip()


# ---------- 需求单 ----------
class TestDemandContract:
    def test_fixed_cells(self):
        ws = openpyxl.load_workbook(TEMPLATE_FILE)["需求单"]
        assert _cell(ws, "A1").startswith("定检中队定检需求单")
        assert "定检运行条件" in _cell(ws, "A2")
        assert "工装" in _cell(ws, "A13") and "航材需求" in _cell(ws, "A13")
        assert _cell(ws, "A14") == "定检专业\n（工具）"
        assert _cell(ws, "A15") == "发动机工具"      # 区标签带"工具"后缀，生成器不得缩写
        assert _cell(ws, "A18") == "机体工具"
        assert _cell(ws, "A21") == "电子工具"
        assert _cell(ws, "A24") == "本次定检新增工作"
        assert _cell(ws, "B24") == "无"
        assert _cell(ws, "A25") == "定检专业\n（航材）"
        assert _cell(ws, "A35") == "三、备用航材需求(定检中队、MCC负责)"
        assert _cell(ws, "A36") == "专业"

    def test_merges_and_print_area(self):
        ws = openpyxl.load_workbook(TEMPLATE_FILE)["需求单"]
        merges = {str(r) for r in ws.merged_cells.ranges}
        for m in ("A1:I1", "A2:I2", "A13:I13", "A35:I35", "B24:I24"):
            assert m in merges, m
        # 已知缺陷（修复计划阶段1）：模板 max_row=42 但打印区只到 I41。
        # 此处锁定现状，阶段 1 修复生成器动态 print_area 后由阶段1用例验证覆盖 42 行。
        pa = str(ws.print_area or "").replace("$", "")
        assert pa.endswith("I41") or "I42" in pa

    def test_data_row_styles(self):
        """数据区参考行样式（动态行应继承而非重设）。"""
        ws = openpyxl.load_workbook(TEMPLATE_FILE)["需求单"]
        assert ws["B15"].font.size == 10.0
        assert ws["B37"].font.size == 14.0


# ---------- 提醒单 ----------
class TestReminderContract:
    def test_fixed_cells(self):
        ws = openpyxl.load_workbook(REMINDER_TEMPLATE_FILE)["工卡提醒"]
        assert _cell(ws, "A1") == "定检工作提醒单"
        assert _cell(ws, "B2") == "定检级别："      # 标签原文，代码只应追加值
        assert "图例" in _cell(ws, "A5") or "红色字体" in _cell(ws, "A5")
        assert _cell(ws, "A6") == "电子"
        assert _cell(ws, "B6") == "发动机"
        assert _cell(ws, "C6") == "机体"

    def test_capacity(self):
        ws = openpyxl.load_workbook(REMINDER_TEMPLATE_FILE)["工卡提醒"]
        assert ws.max_row >= 5000


# ---------- 工具清单 ----------
class TestToolsContract:
    def test_zones_and_preprinted(self):
        ws = openpyxl.load_workbook(TOOLS_TEMPLATE_FILE)["定检工具"]
        merges = {str(r) for r in ws.merged_cells.ranges}
        for m in ("A1:G1", "A3:A18", "A19:A28", "A29:A38", "A39:A48"):
            assert m in merges, m
        assert _cell(ws, "B3") == "手套"
        assert _cell(ws, "B4") == "毛巾"
        for first, last in ((19, 28), (29, 38), (39, 48)):
            for col in ("B", "D"):
                assert ws[f"{col}{first}"].border.top.style == "medium", f"{col}{first}"
                assert ws[f"{col}{last}"].border.bottom.style == "medium", f"{col}{last}"

    def test_no_internal_merges_in_zones(self):
        """区内除 A 标签外无内部合并（_extend_zone 只平移下方合并的假设）。"""
        ws = openpyxl.load_workbook(TOOLS_TEMPLATE_FILE)["定检工具"]
        zones = [(3, 18), (19, 28), (29, 38), (39, 48)]
        for rng in ws.merged_cells.ranges:
            if rng.min_col == 1 and rng.max_col == 1:
                continue   # A 列区标签
            for start, end in zones:
                assert not (start <= rng.min_row and rng.max_row <= end), \
                    f"区内出现内部合并 {rng}"


# ---------- 航化清单 ----------
class TestMaterialsContract:
    def test_zones_and_preprinted(self):
        ws = openpyxl.load_workbook(MATERIALS_TEMPLATE_FILE)["定检航化"]
        merges = {str(r) for r in ws.merged_cells.ranges}
        for m in ("A1:H1", "A3:A11", "A12:A17", "A18:A24"):
            assert m in merges, m
        assert "纸胶带" in _cell(ws, "B3")
        assert _cell(ws, "B12") == "清洗剂"       # 外场预印
        assert "警戒" in _cell(ws, "A25")
        assert _cell(ws, "B26") == "无纺布"
        assert ws["C18"].font.size == 12.0        # 非例行区字号差异（不得被统一覆盖）


# ---------- 改版清单 ----------
class TestCheckContract:
    def test_fixed_cells_and_capacity(self):
        ws = openpyxl.load_workbook(CHECK_TEMPLATE_FILE)["改版清单"]
        assert _cell(ws, "A1") == "工卡改版清单"
        merges = {str(r) for r in ws.merged_cells.ranges}
        assert "A1:C1" in merges
        assert (_cell(ws, "A2"), _cell(ws, "B2"), _cell(ws, "C2")) == ("电子", "发动机", "机体")
        assert ws.max_row >= 4996
        assert ws["A3"].fill.fill_type == "solid"   # 数据区预置绿底


# ---------- 生成器对固定区的保护（隔离副本验证行为） ----------
class TestGeneratorRespectsFixedZone:
    def test_zone_label_missing_raises(self, tmp_path, monkeypatch):
        """模板结构异常必须明确报错（区标签被改坏的场景）。"""
        import reqman.services.checklist_generator as cg

        dst = tmp_path / "tools_broken.xlsx"
        wb = openpyxl.load_workbook(TOOLS_TEMPLATE_FILE)
        wb["定检工具"]["A3"] = "通用工具区"   # 改坏区标签
        wb.save(dst)
        monkeypatch.setattr(cg, "TOOLS_TEMPLATE_FILE", dst)
        with pytest.raises(RuntimeError, match="模板结构异常"):
            cg.generate_tool_list(
                {"reg": "B-1234", "description": "T", "date": "2026-09-10"},
                [{"device_name": "扳手", "part_number": "PN", "quantity": "1",
                  "category": "发动机", "usage_type": "必须使用"}])
