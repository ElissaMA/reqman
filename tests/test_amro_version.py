"""T6/T7 工卡版本域：全库检查、作废移库、改版清单 Excel、提醒单附加区块（全 mock）"""
import asyncio
import io
import re
import zipfile

import openpyxl
import pytest

from reqman.models.cancelled_card_store import CancelledCardStore
from reqman.services import amro_sync


def _run(coro):
    return asyncio.run(coro)


def _jcrow(jcno, wd, task="RST", zy="电子", title="检查救生衣"):
    """TD_JC_*_LIST 行（WRITE_DATE 为版本主轴，两清单零缺失）"""
    return {"JC_NO": jcno, "WRITE_DATE": wd, "TASK": task, "ZY": zy, "JCTITLE": title}


@pytest.fixture
def fake_amro_cards():
    """SMJC：库内例行卡（含改版 1 张）；非 SMJC 插件一律返回空（EO 卡走实体直查 mock）"""
    async def fetch(client, cookies, plugin, base_form, **kw):
        if plugin == "TD_JC_SMJC_LIST":
            return [_jcrow("CSCA320-256652-01-1-X", "2026-08-01 09:00:00")]
        return []
    return fetch


@pytest.fixture
def fake_entity_query():
    """实体端点直查 mock：两张在册 EO 卡返回实体行，其余查无（→作废）。"""
    async def query(client, cookies, plugin, form, **kw):
        codes = {"EOJC-A320-31-2026-007-A": "2026-07-15 14:00:00",
                 "EOJC-A320-99-2026-999-Z": "2026-07-20 08:00:00"}
        jcno = form.get("jcno")
        if jcno in codes:
            return {"code": 200, "data": {"JC_NO": jcno, "WRITE_DATE": codes[jcno]}}
        return {"code": 200, "data": {}}
    return query


class TestFullVersionCheck:
    def test_revised_updates_card_and_logs(self, json_store, fake_amro_cards, monkeypatch):
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST", "")
        rep = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fake_amro_cards))
        card = json_store.get(r["id"])
        assert card["write_date"] == "2026-08-01 09:00:00"
        # 原库无编写日期、本次被填入 → 归入「新增」而非「改版」
        assert len(rep["revised"]) == 0
        assert rep["new_added"][0]["old_wd"] == "" and rep["new_added"][0]["new_wd"] == "2026-08-01 09:00:00"
        assert any(c["field"] == "write_date" for l in json_store._read()["card_logs"]
                   for c in l["changes"])

    def test_cancelled_report_only_without_store(self, json_store, fake_amro_cards,
                                                 fake_entity_query, monkeypatch):
        """未传作废库时保持旧行为：作废只入报告不删卡（向后兼容）。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("EOJC-A320-57-2025-002-B", "卡", "机体", "", "")
        rep = _run(amro_sync.full_version_check(json_store, None, {},
                                                fetch=fake_amro_cards, query=fake_entity_query))
        assert "EOJC-A320-57-2025-002-B" in [c["task_code"] for c in rep["cancelled"]]
        assert json_store.get(r["id"]) is not None   # 仅报告不删除

    def test_cancelled_moved_to_library(self, json_store, fake_amro_cards, fake_entity_query,
                                        monkeypatch, tmp_path):
        """作废整卡自动移入作废工卡库：主库删除、承接全部原字段+作废元数据，报告照旧。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        cancelled_store = CancelledCardStore(str(tmp_path / "cancelled_cards.json"))
        r = json_store.add("EOJC-A320-57-2025-002-B", "旧卡", "机体", "EO", "",
                           reminder_type="重点提醒")
        json_store.update(r["id"], write_date="2026-01-01", tools_confirmed=True)
        rep = _run(amro_sync.full_version_check(json_store, None, {},
                                                fetch=fake_amro_cards, query=fake_entity_query,
                                                cancelled_store=cancelled_store))
        assert json_store.get(r["id"]) is None       # 主库已删除
        assert json_store.find_by_code("EOJC-A320-57-2025-002-B") is None
        rec = cancelled_store.find_by_code("EOJC-A320-57-2025-002-B")
        assert rec["task_name"] == "旧卡" and rec["category"] == "机体"
        assert rec["orig_id"] == r["id"]
        assert rec["cancel_source"] == "full_version"
        assert rec["cancelled_at"]
        assert rec["reminder_type"] == "重点提醒"     # 承接原卡全部信息
        assert rec["tools_confirmed"] is True and rec["write_date"] == "2026-01-01"
        assert rep["cancelled"][0]["task_code"] == "EOJC-A320-57-2025-002-B"   # 报告照旧

    def test_checked_count_excludes_dp(self, json_store, fake_amro_cards, fake_entity_query,
                                       monkeypatch, tmp_path):
        """checked = 参与比对的非 DP 卡数（DP 跳过不计数）；不再调用 EOJC 深分页。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        cancelled_store = CancelledCardStore(str(tmp_path / "cancelled_cards.json"))
        json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST", "")
        json_store.add("EOJC-A320-57-2025-002-B", "旧卡", "机体", "EO", "")
        json_store.add("DP1000000739", "DP项目卡", "机体", "", "")
        plugins = []

        async def fetch(client, cookies, plugin, base_form, **kw):
            plugins.append(plugin)
            return await fake_amro_cards(client, cookies, plugin, base_form, **kw)

        rep = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fetch,
                                                query=fake_entity_query,
                                                cancelled_store=cancelled_store))
        assert rep["checked"] == 2
        assert "TD_JC_ALL_EOJC_LIST" not in plugins   # EOJC 深分页已根除

    def test_unchanged_card_not_relogged(self, json_store, fake_amro_cards, monkeypatch):
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST", "")
        json_store.update(r["id"], write_date="2026-08-01 09:00:00")   # 已是最新
        logs_before = len(json_store.get_version_logs())
        rep = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fake_amro_cards))
        assert rep["revised"] == []
        assert len(json_store.get_version_logs()) == logs_before   # 不重复记版本日志

    def test_same_date_different_time_not_relogged(self, json_store, fake_amro_cards, monkeypatch):
        """日期部分比对：界面 date 只存 YYYY-MM-DD，AMRO 同日不同时间不再误报版本变动。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST", "")
        json_store.update(r["id"], write_date="2026-08-01")   # 手动录入仅日期
        logs_before = len(json_store.get_version_logs())
        rep = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fake_amro_cards))
        assert rep["revised"] == []
        assert len(json_store.get_version_logs()) == logs_before

        async def fetch_revised(client, cookies, plugin, base_form, **kw):  # 真正跨日期改版
            if plugin == "TD_JC_SMJC_LIST":
                return [_jcrow("CSCA320-256652-01-1-X", "2026-08-05 09:00:00")]
            return []
        rep2 = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fetch_revised))
        assert len(rep2["revised"]) == 1
        assert json_store.get(r["id"])["write_date"] == "2026-08-05 09:00:00"


    def test_dp_cards_skipped(self, json_store, fake_amro_cards, monkeypatch):
        """DP 开头工卡（DP 项目）不在 AMRO 清单体系内：全库检查跳过不比对、不报作废。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("DP1000000739", "DP项目卡", "机体", "", "")
        rep = _run(amro_sync.full_version_check(json_store, None, {}, fetch=fake_amro_cards))
        assert all(c["task_code"] != "DP1000000739"
                   for c in rep["revised"] + rep["cancelled"])
        assert json_store.get(r["id"]) is not None


class TestVersionReportExcel:
    def test_dot_date(self):
        assert amro_sync._dot_date("20260901_143025") == "2026.09.01"
        assert amro_sync._dot_date("") == ""
        assert amro_sync._dot_date("bad") == ""

    def test_package_report_label(self):
        """版本报告标识：机号+描述+开工日期，日期归一化点分、可缺项。"""
        assert amro_sync.package_report_label(
            {"aircraft_info": {"reg": "B-1234", "description": "46A", "date": "2026-09-05"}}
        ) == "B-1234 46A 2026.09.05"
        assert amro_sync.package_report_label(
            {"reg": "B-1", "description": "D", "date": "2026.09.05"}) == "B-1 D 2026.09.05"
        assert amro_sync.package_report_label({"reg": "B-1", "description": "D"}) == "B-1 D"
        assert amro_sync.package_report_label({}) == ""

    def test_build_package_label_unifies_display_and_report(self):
        """build_package_label 统一展示/报告标识；with_date 控制是否附加点分日期。"""
        pkg = {"aircraft_info": {"reg": "B-1234", "description": "46A", "date": "2026-09-05"}}
        assert amro_sync.build_package_label(pkg) == "B-1234 46A"
        assert amro_sync.build_package_label(pkg, with_date=True) == "B-1234 46A 2026.09.05"
        # 旧别名等价
        assert amro_sync.package_display_label(pkg) == amro_sync.build_package_label(pkg)
        assert amro_sync.package_report_label(pkg) == amro_sync.build_package_label(pkg, with_date=True)

    def test_report_excel_reminder_template_layout(self):
        """改版清单以专用模板输出：标题含「查询日期」、单单元格三行、第三行按类型着色、保留绿底。"""
        buf = amro_sync.build_version_report_excel({
            "revised": [
                {"task_code": "C-1", "task_name": "卡一", "category": "电子",
                 "old_wd": "2026-07-01 09:00:00", "new_wd": "2026-08-01 09:00:00"},
            ],
            "new_added": [
                {"task_code": "C-2", "task_name": "卡二", "category": "发动机",
                 "old_wd": "", "new_wd": "2026-08-01 09:00:00"},
            ],
            "cancelled": [
                {"task_code": "C-3", "task_name": "卡三", "category": "机体"},
                {"task_code": "C-4", "task_name": "特检卡", "category": "特检"},
            ],
        }, title_label="B-1234 46A 2026.09.05", finished_date="2026.09.01")
        wb = openpyxl.load_workbook(io.BytesIO(buf))
        assert wb.sheetnames == ["改版清单"]
        ws = wb["改版清单"]
        assert ws["A1"].value == "工卡改版清单（B-1234 46A 2026.09.05）查询日期2026.09.01"
        assert (ws["A2"].value, ws["B2"].value, ws["C2"].value) == ("电子", "发动机", "机体")
        assert ws["A2"].fill.start_color.rgb == "FF00703C"   # 表头深绿白字样式保留
        a3 = ws["A3"]   # 电子列：改版卡一，单单元格三行
        assert a3.value == "C-1\n卡一\n2026-07-01→2026-08-01"
        assert a3.fill.start_color.rgb == "FFAAD296"         # 保留原绿底
        assert ws["B3"].value == "C-2\n卡二\n新增 2026-08-01"  # 原库无日期 → 新增（第三行标红）
        assert ws["B3"].fill.start_color.rgb == "FFC8DCB4"
        c3 = ws["C3"]   # 机体列：作废卡三
        assert c3.value == "C-3\n卡三\n作废"
        assert ws["D2"].value is None and ws["D3"].value is None   # 特检不输出
        assert sorted(str(r) for r in ws.merged_cells.ranges) == ["A1:C1"]   # 仅标题合并
        assert ws["A5"].value is None and ws["B5"].value is None and ws["C5"].value is None

    def test_report_excel_third_line_color_by_type(self, tmp_path):
        """改版清单：单元格仅第三行（标记行）着色，工卡号/名称保持黑字；
        第三行按类型配色：新增=红、改版=蓝、作废=黑。"""
        buf = amro_sync.build_version_report_excel({
            "revised": [
                {"task_code": "R-1", "task_name": "改版卡", "category": "电子",
                 "old_wd": "2026-07-01", "new_wd": "2026-08-01 09:00:00"},
            ],
            "new_added": [
                {"task_code": "N-1", "task_name": "新增卡", "category": "机体",
                 "old_wd": "", "new_wd": "2026-09-01 10:00:00"},
            ],
            "cancelled": [
                {"task_code": "X-1", "task_name": "作废卡", "category": "发动机"},
            ],
        })
        p = tmp_path / "rt_red.xlsx"
        p.write_bytes(buf)
        with zipfile.ZipFile(p) as z:
            xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8", "ignore")

        def runs():
            for m in re.finditer(r"<r>(.*?)</r>", xml, re.DOTALL):
                r = m.group(1)
                t = re.search(r"<t[^>]*>(.*?)</t>", r, re.DOTALL)
                c = re.search(r'<color rgb="([0-9A-Fa-f]+)"', r)
                yield (t.group(1) if t else "", c.group(1).upper() if c else None)

        rs = list(runs())
        red = [t for t, col in rs if col == "FFFF0000"]
        blue = [t for t, col in rs if col == "FF0000FF"]
        black = [t for t, col in rs if col == "FF000000"]
        # 第三行（标记行）按类型着色：新增=红、改版=蓝、作废=黑
        assert any(t.startswith("新增") for t in red), "新增标记应为红"
        assert any("→" in t for t in blue), "改版标记应为蓝"
        assert any(t == "作废" for t in black), "作废标记应为黑"
        # 工卡号 / 名称行保持黑字
        assert {"R-1", "N-1", "X-1"}.issubset(set(black)), "工卡号应为黑字"
        assert any("改版卡" in t for t in black) and any("新增卡" in t for t in black)


class TestCheckCardsAgainstAmro:
    """包级版本检查（工作包行级查询）：条目含 task_name/category，无 new_by_category"""

    def test_dp_codes_excluded(self, json_store, monkeypatch):
        """DP 开头工卡不参与包级版本检查：不发查询、不误报作废。"""
        json_store.add("DP1000000739", "DP项目卡", "机体", "", "")
        json_store.add("EOJC-A320-57-2025-002-B", "旧卡", "机体", "EO", "")
        queried = []

        async def query(client, cookies, plugin, form, **kw):
            queried.append(form.get("jcno"))
            return {"code": 200, "data": {}}   # 查无 → 作废

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["DP1000000739", "EOJC-A320-57-2025-002-B"], query=query))
        assert queried == ["EOJC-A320-57-2025-002-B"]   # DP 不发任何查询
        assert [c["task_code"] for c in rep["cancelled"]] == ["EOJC-A320-57-2025-002-B"]
        assert all(c["task_code"] != "DP1000000739"
                   for c in rep["revised"] + rep["cancelled"])

    def test_revised_and_cancelled_with_fields(self, json_store, fake_amro_cards, monkeypatch):
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r_csca = json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST")
        json_store.update(r_csca["id"], write_date="2026-07-01 09:00:00")
        json_store.add("EOJC-A320-57-2025-002-B", "旧卡", "机体", "EO", "")

        async def query(client, cookies, plugin, form, **kw):   # EO 卡逐条直查
            if form.get("jcno") == "EOJC-A320-99-2026-999-Z":
                return {"code": 200, "data": {"JC_NO": "EOJC-A320-99-2026-999-Z",
                                              "WRITE_DATE": "2026-07-20 08:00:00"}}
            return {"code": 200, "data": {}}   # 57-2025-002-B 查无 → 作废

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {},
            ["CSCA320-256652-01-1-X", "EOJC-A320-57-2025-002-B",
             "EOJC-A320-99-2026-999-Z"],   # 库内无此卡（AMRO 有）→ 跳过
            fetch=fake_amro_cards, query=query))
        assert "new_by_category" not in rep
        rev = rep["revised"][0]
        assert rev["task_code"] == "CSCA320-256652-01-1-X"
        assert rev["task_name"] == "检查救生衣" and rev["category"] == "电子"
        can = rep["cancelled"][0]
        assert can["task_code"] == "EOJC-A320-57-2025-002-B"
        assert can["task_name"] == "旧卡" and can["category"] == "机体"
        codes = {r["task_code"] for r in rep["revised"] + rep["cancelled"]}
        assert "EOJC-A320-99-2026-999-Z" not in codes

    def test_cancelled_moved_to_library(self, json_store, monkeypatch, tmp_path):
        """包级检查作废 → 整卡移入作废工卡库；checked = 工作包内去重卡号数（DP 排除）。"""
        cancelled_store = CancelledCardStore(str(tmp_path / "cancelled_cards.json"))
        json_store.add("EOJC-A320-57-2025-002-B", "旧卡", "机体", "EO", "")

        async def query(client, cookies, plugin, form, **kw):
            return {"code": 200, "data": {}}   # 查无 → 作废

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["EOJC-A320-57-2025-002-B", "DP1000000739"],
            query=query, cancelled_store=cancelled_store))
        assert [c["task_code"] for c in rep["cancelled"]] == ["EOJC-A320-57-2025-002-B"]
        assert rep["checked"] == 1
        assert json_store.find_by_code("EOJC-A320-57-2025-002-B") is None
        rec = cancelled_store.find_by_code("EOJC-A320-57-2025-002-B")
        assert rec["cancel_source"] == "package_version"


class TestVersionSummaryText:
    def test_base_format(self):
        """两处版本检查共用的基础摘要文案：改版X张，新增N张，作废Y张，共检查Z张。"""
        assert amro_sync._version_summary(2, 1, 10) == "改版 2 张，新增 0 张，作废 1 张，共检查 10 张"
        assert amro_sync._version_summary(0, 0, 0) == "改版 0 张，新增 0 张，作废 0 张，共检查 0 张"
        assert amro_sync._version_summary(2, 1, 10, 3) == "改版 2 张，新增 3 张，作废 1 张，共检查 10 张"


class TestVersionPullSplit:
    """v3.6.0 版本查询提速：SMJC/EOJC 机队筛选 + 例行卡包跳过 EOJC 深分页 + EO 逐卡直查"""

    def test_pull_smjc_versions_fleet_filter(self, monkeypatch):
        """SMJC 表单带 fleet=A320（2026-09-01 实测 1319→745）；EOJC 深分页已移除。"""
        seen = {}

        async def fetch(client, cookies, plugin, base_form, **kw):
            seen[plugin] = dict(base_form)
            return [_jcrow("CSCA320-256652-01-1-X", "2026-08-01 09:00:00")]

        rows = _run(amro_sync._pull_smjc_versions(None, {}, fetch=fetch))
        assert seen == {"TD_JC_SMJC_LIST": {"status": "ISSUED", "jcStatus": "Y",
                                            "rows": 500, "fleet": "A320"}}
        assert list(rows) == ["CSCA320-256652-01-1-X"]

    def test_get_entity_by_jcno_returns_row_or_none(self, monkeypatch):
        """实体端点：data 单对象 → 行 dict；空 data → None。"""
        seen = {}

        async def query(client, cookies, plugin, form, **kw):
            seen[plugin] = dict(form)
            if form.get("jcno") == "HAS":
                return {"code": 200, "data": {"JC_NO": "HAS", "WRITE_DATE": "2026-08-01 09:00:00"}}
            return {"code": 200, "data": {}}

        row = _run(amro_sync._get_entity_by_jcno(None, {}, "HAS", query=query))
        assert row == {"JC_NO": "HAS", "WRITE_DATE": "2026-08-01 09:00:00"}
        assert seen == {"TD_JC_ALL_GET_ENTITY_BY_JCNO": {"jcno": "HAS"}}
        assert _run(amro_sync._get_entity_by_jcno(None, {}, "MISS", query=query)) is None

    def test_csca_only_package_skips_eojc(self, json_store, monkeypatch):
        """包内全为定检例行卡（CSCA 前缀）→ 只拉 SMJC，跳过 EOJC 深分页（秒级返回）。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r_csca = json_store.add("CSCA320-256652-01-1-X", "检查救生衣", "电子", "RST")
        json_store.update(r_csca["id"], write_date="2026-07-01 09:00:00")
        called = []

        async def fetch(client, cookies, plugin, base_form, **kw):
            called.append(plugin)
            if plugin == "TD_JC_SMJC_LIST":
                return [_jcrow("CSCA320-256652-01-1-X", "2026-08-01 09:00:00")]
            return []

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["CSCA320-256652-01-1-X"], fetch=fetch))
        assert called == ["TD_JC_SMJC_LIST"]
        assert len(rep["revised"]) == 1

    def test_eo_package_uses_per_jcno_query(self, json_store, monkeypatch):
        """含 EO 卡 → SMJC 全量拉 + EO 逐个按卡号直查（不再全量拉 EOJC 深分页）。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r_eo = json_store.add("EOJC-A320-31-2026-007-A", "改装卡", "电子", "EO")
        json_store.update(r_eo["id"], write_date="2026-07-01 09:00:00")
        fetched, queried = [], []

        async def fetch(client, cookies, plugin, base_form, **kw):
            fetched.append((plugin, dict(base_form)))
            if plugin == "TD_JC_SMJC_LIST":
                return [_jcrow("CSCA320-256652-01-1-X", "2026-08-01 09:00:00")]
            return []

        async def query(client, cookies, plugin, form, **kw):
            queried.append((plugin, dict(form)))
            return {"code": 200, "data": {"JC_NO": "EOJC-A320-31-2026-007-A",
                                          "WRITE_DATE": "2026-07-15 14:00:00"}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {},
            ["CSCA320-256652-01-1-X", "EOJC-A320-31-2026-007-A"],
            fetch=fetch, query=query))
        assert [p for p, _ in fetched] == ["TD_JC_SMJC_LIST"]          # 无 EOJC 全量拉取
        assert queried == [("TD_JC_ALL_GET_ENTITY_BY_JCNO",
                            {"jcno": "EOJC-A320-31-2026-007-A"})]      # EO 逐卡直查
        assert [r["task_code"] for r in rep["revised"]] == ["EOJC-A320-31-2026-007-A"]


class TestFamilyOf:
    def test_four_families(self):
        assert amro_sync._family_of("CSCA320-256652-01-1-X") == "CSC"
        assert amro_sync._family_of("FLA-A320-1001-1") == "FLA"
        assert amro_sync._family_of("EOJC-A320-31-2026-007-A") == "EO"
        assert amro_sync._family_of("QECJC-A320-1001-1") == "QEC"
        assert amro_sync._family_of("ERJC-A320-1001-1") == "QEC"

    def test_out_of_scope(self):
        assert amro_sync._family_of("DP1000000739") == ""
        assert amro_sync._family_of("NRC12345") == ""
        assert amro_sync._family_of("LS123") == ""
        assert amro_sync._family_of("") == ""
        assert amro_sync._is_in_scope("NRC12345") is False
        assert amro_sync._is_in_scope("CSCA320-256652-01-1-X") is True


class TestPullAmroFamily:
    def test_prefix_filter_and_fleet(self, monkeypatch):
        """全量拉取按前缀过滤（邻族不并入）+ 透传 fleet=A320。"""
        seen = {}

        async def fetch(client, cookies, plugin, base_form, **kw):
            seen[plugin] = dict(base_form)
            return [
                _jcrow("FLA-A320-1001-1", "2026-08-01 09:00:00"),
                _jcrow("CSCA320-256652-01-1-X", "2026-08-01 09:00:00"),  # 邻族应被过滤
            ]

        rows = _run(amro_sync._pull_amro_family(
            None, {}, "TD_JC_NRCJC_LIST",
            {"rows": 500, "fleet": "A320"}, ("FLA",), fetch=fetch))
        assert seen["TD_JC_NRCJC_LIST"] == {"rows": 500, "fleet": "A320"}
        assert list(rows) == ["FLA-A320-1001-1"]


class TestCollectFourFamilies:
    def test_qec_full_pull_via_qecjc_list(self, json_store, monkeypatch):
        """QEC/ER（QECJC*/ERJC*）走专用 TD_JC_ALL_QECJC_LIST 全量拉（fleet=A320），不逐卡。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r1 = json_store.add("QECJC-A320-1001-1", "卡", "机体", "QEC", "")
        json_store.update(r1["id"], write_date="2026-07-01 09:00:00")
        r2 = json_store.add("ERJC-A320-1001-1", "卡", "机体", "QEC", "")
        json_store.update(r2["id"], write_date="2026-07-01 09:00:00")
        fetched, queried = [], []

        async def fetch(client, cookies, plugin, base_form, **kw):
            fetched.append((plugin, dict(base_form)))
            if plugin == "TD_JC_ALL_QECJC_LIST":
                return [_jcrow("QECJC-A320-1001-1", "2026-08-01 09:00:00"),
                        _jcrow("ERJC-A320-1001-1", "2026-08-02 09:00:00")]
            return []

        async def query(client, cookies, plugin, form, **kw):
            queried.append(plugin)
            return {"code": 200, "data": {}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["QECJC-A320-1001-1", "ERJC-A320-1001-1"],
            fetch=fetch, query=query))
        assert [p for p, _ in fetched] == ["TD_JC_ALL_QECJC_LIST"]
        assert queried == []   # 不逐卡
        assert {r["task_code"] for r in rep["revised"]} == {"QECJC-A320-1001-1", "ERJC-A320-1001-1"}

    def test_fla_full_pull_via_nrcjc_list_when_confirmed(self, json_store, monkeypatch):
        """FLA 端点已确认加入白名单后，走 TD_JC_NRCJC_LIST 全量拉；不逐卡。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("FLA-A320-1001-1", "卡", "机体", "FLA", "")
        json_store.update(r["id"], write_date="2026-07-01 09:00:00")
        fetched, queried = [], []

        async def fetch(client, cookies, plugin, base_form, **kw):
            fetched.append((plugin, dict(base_form)))
            if plugin == "TD_JC_NRCJC_LIST":
                return [_jcrow("FLA-A320-1001-1", "2026-08-01 09:00:00")]
            return []

        async def query(client, cookies, plugin, form, **kw):
            queried.append((plugin, dict(form)))
            return {"code": 200, "data": {}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["FLA-A320-1001-1"], fetch=fetch, query=query))
        assert [p for p, _ in fetched] == ["TD_JC_NRCJC_LIST"]
        assert queried == []   # 不逐卡
        assert [r["task_code"] for r in rep["revised"]] == ["FLA-A320-1001-1"]

    def test_fla_falls_back_to_per_card_when_whitelist_excludes_it(self, json_store, monkeypatch):
        """防御性：若 FLA 端点不在白名单，回退逐卡直查；不调用未确认全量端点。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        monkeypatch.setattr(amro_sync.amro, "READONLY_PLUGINS",
                            amro_sync.amro.READONLY_PLUGINS - {"TD_JC_NRCJC_LIST"})
        r = json_store.add("FLA-A320-1001-1", "卡", "机体", "FLA", "")
        json_store.update(r["id"], write_date="2026-07-01 09:00:00")
        fetched, queried = [], []

        async def fetch(client, cookies, plugin, base_form, **kw):
            fetched.append(plugin)
            return []

        async def query(client, cookies, plugin, form, **kw):
            queried.append((plugin, dict(form)))
            if form.get("jcno") == "FLA-A320-1001-1":
                return {"code": 200, "data": {"JC_NO": "FLA-A320-1001-1",
                                              "WRITE_DATE": "2026-08-01 09:00:00"}}
            return {"code": 200, "data": {}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["FLA-A320-1001-1"], fetch=fetch, query=query))
        assert "TD_JC_NRCJC_LIST" not in fetched   # 未确认端点不被调用
        assert queried == [("TD_JC_ALL_GET_ENTITY_BY_JCNO", {"jcno": "FLA-A320-1001-1"})]
        assert [r["task_code"] for r in rep["revised"]] == ["FLA-A320-1001-1"]

    def test_qec_full_pull_failure_falls_back_per_card(self, json_store, monkeypatch):
        """QEC/ER 全量拉取失败 → 回退逐卡直查，避免误判作废。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        r = json_store.add("QECJC-A320-1001-1", "卡", "机体", "QEC", "")
        json_store.update(r["id"], write_date="2026-07-01 09:00:00")
        fetched, queried = [], []

        async def fetch(client, cookies, plugin, base_form, **kw):
            fetched.append(plugin)
            raise RuntimeError("AMRO 全量拉取失败（瞬态）")

        async def query(client, cookies, plugin, form, **kw):
            queried.append((plugin, dict(form)))
            if form.get("jcno") == "QECJC-A320-1001-1":
                return {"code": 200, "data": {"JC_NO": "QECJC-A320-1001-1",
                                              "WRITE_DATE": "2026-08-01 09:00:00"}}
            return {"code": 200, "data": {}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["QECJC-A320-1001-1"], fetch=fetch, query=query))
        assert "TD_JC_ALL_QECJC_LIST" in fetched
        assert queried == [("TD_JC_ALL_GET_ENTITY_BY_JCNO", {"jcno": "QECJC-A320-1001-1"})]
        assert [r["task_code"] for r in rep["revised"]] == ["QECJC-A320-1001-1"]

    def test_out_of_scope_skipped_not_cancelled(self, json_store, monkeypatch, tmp_path):
        """非四家族卡（如 NRC/航线）全库检查跳过：不查、不误报作废、保留在主库。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        cancelled_store = CancelledCardStore(str(tmp_path / "cancelled.json"))
        r = json_store.add("NRC12345", "非四家族卡", "机体", "", "")

        async def fetch(client, cookies, plugin, base_form, **kw):
            return []

        rep = _run(amro_sync.full_version_check(
            json_store, None, {}, fetch=fetch, cancelled_store=cancelled_store))
        assert rep["checked"] == 0
        assert rep["cancelled"] == []
        assert json_store.get(r["id"]) is not None   # 保留在主库

    def test_package_out_of_scope_skipped(self, json_store, monkeypatch):
        """包内非四家族卡不参与版本检查：checked 不计入、不查、不报作废。"""
        monkeypatch.setattr(amro_sync.amro.time, "monotonic", lambda: 1e9)
        monkeypatch.setattr(amro_sync.amro.time, "sleep", lambda s: None)
        json_store.add("NRC12345", "非四家族卡", "机体", "", "")
        queried = []

        async def query(client, cookies, plugin, form, **kw):
            queried.append(form.get("jcno"))
            return {"code": 200, "data": {}}

        rep = _run(amro_sync.check_cards_against_amro(
            json_store, None, {}, ["NRC12345"], query=query))
        assert queried == []           # 非四家族不查
        assert rep["checked"] == 0
        assert rep["cancelled"] == []
