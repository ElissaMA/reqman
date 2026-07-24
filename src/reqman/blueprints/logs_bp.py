"""日志蓝图 —— 操作日志的浏览、导出与删除"""

import logging
from datetime import datetime, timedelta
from urllib.parse import urlencode

from flask import Blueprint, render_template, request, jsonify, flash, redirect, send_file
from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO

from ..utils.response import api_success, api_error
from ..utils.error_handlers import _wants_json, raise_or_flash, NotFoundError, ServerError

logger = logging.getLogger(__name__)

bp = Blueprint("logs", __name__, url_prefix="/card")

# 操作类型中文显示
OPERATION_LABELS = {
    "add": "新增",
    "update": "修改",
    "delete": "删除",
}

# 目标类型中文显示
TARGET_TYPE_LABELS = {
    "card": "工卡",
    "set": "工卡组",
    "aircraft": "飞机信息",
}

# 字段中文标签（用于展示变更内容）
FIELD_LABELS = {
    "task_code": "工卡编码",
    "task_name": "工卡名称",
    "category": "分类",
    "work_area": "工作区域",
    "aircraft_type": "机型",
    "tools": "工具",
    "materials": "航材",
    "name": "名称",
    "set_name": "组名称",
    "description": "描述",
    "note": "备注",
    "status": "状态",
    "card_codes": "工卡列表",
    "interval_type": "间隔类型",
    "interval_value": "间隔值",
}


def _get_service():
    """获取 CardService 实例"""
    return current_app.extensions.get("card_service")


def _format_changes(changes: list) -> str:
    """将 changes 列表转为可读字符串"""
    if not changes:
        return ""
    parts = []
    for c in changes:
        field = c.get("field", "")
        label = FIELD_LABELS.get(field, field)
        old_val = c.get("old", "")
        new_val = c.get("new", "")
        if isinstance(old_val, list):
            old_val = f"[{len(old_val)} 项]"
        if isinstance(new_val, list):
            new_val = f"[{len(new_val)} 项]"
        parts.append(f"{label}: {old_val} → {new_val}")
    return "; ".join(parts)


def _build_filtered_logs(args: dict) -> list:
    """根据筛选条件返回日志列表"""
    service = _get_service()
    if not service:
        return []
    logs = list(service.store.get_logs())

    # 操作类型筛选
    operation = args.get("operation", "").strip()
    if operation:
        logs = [l for l in logs if l.get("operation") == operation]

    # 目标类型筛选
    target_type = args.get("target_type", "").strip()
    if target_type:
        logs = [l for l in logs if l.get("target_type") == target_type]

    # 关键词搜索
    keyword = args.get("keyword", "").strip()
    if keyword:
        kw = keyword.lower()
        logs = [l for l in logs if
                kw in (l.get("target_identifier") or "").lower() or
                kw in (l.get("target_name") or "").lower()]

    # 时间范围筛选
    date_from = args.get("date_from", "").strip()
    date_to = args.get("date_to", "").strip()
    if date_from:
        logs = [l for l in logs if (l.get("timestamp") or "") >= date_from]
    if date_to:
        # 包含当天
        logs = [l for l in logs if (l.get("timestamp") or "") <= date_to + "T23:59:59"]

    # 按时间降序排列
    logs.sort(key=lambda l: l.get("timestamp", ""), reverse=True)
    return logs


# ===================== 路由 =====================


@bp.route("/logs")
def list_logs():
    """日志列表页面"""
    logs = _build_filtered_logs(request.args)

    page = request.args.get("page", 1, type=int)
    per_page = 20
    total = len(logs)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    start = (page - 1) * per_page
    end = start + per_page
    page_logs = logs[start:end]

    # 保留筛选参数（除 page 外）用于分页链接
    qp = {k: v for k, v in request.args.items() if k != "page"}
    query_params = "&" + urlencode(qp) if qp else ""

    return render_template(
        "cards/logs.html",
        logs=page_logs,
        page=page,
        total_pages=total_pages,
        total=total,
        query_params=query_params,
        operation_labels=OPERATION_LABELS,
        target_type_labels=TARGET_TYPE_LABELS,
        field_labels=FIELD_LABELS,
        format_changes=_format_changes,
    )


@bp.route("/logs/export", methods=["POST"])
def export_logs():
    """导出选中日志为 Excel"""
    log_ids = request.form.getlist("log_ids[]")

    service = _get_service()
    if not service:
        if _wants_json():
            return api_error("服务不可用", "SERVICE_UNAVAILABLE", 500)
        flash("服务不可用", "error")
        return redirect(request.headers.get("Referer", "/"))

    all_logs = list(service.store.get_logs())

    if log_ids:
        ids = set()
        for lid in log_ids:
            try:
                ids.add(int(lid))
            except (ValueError, TypeError):
                pass
        selected = [l for l in all_logs if l.get("id") in ids]
    else:
        selected = list(all_logs)

    selected.sort(key=lambda l: l.get("timestamp", ""), reverse=True)

    # 生成 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "操作日志"

    # 表头样式
    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="1A3A5C", end_color="1A3A5C", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["ID", "时间", "操作", "目标类型", "目标标识", "目标名称", "变更内容"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 数据行
    data_alignment = Alignment(vertical="center", wrap_text=True)
    for row_idx, log in enumerate(selected, 2):
        ws.cell(row=row_idx, column=1, value=log.get("id", "")).border = thin_border
        ts = log.get("timestamp", "")
        if isinstance(ts, str) and len(ts) >= 10:
            ts = ts[:10]  # 只显示日期
        ws.cell(row=row_idx, column=2, value=ts).border = thin_border
        op = log.get("operation", "")
        ws.cell(row=row_idx, column=3, value=OPERATION_LABELS.get(op, op)).border = thin_border
        tt = log.get("target_type", "")
        ws.cell(row=row_idx, column=4, value=TARGET_TYPE_LABELS.get(tt, tt)).border = thin_border
        ws.cell(row=row_idx, column=5, value=log.get("target_identifier", "")).border = thin_border
        ws.cell(row=row_idx, column=6, value=log.get("target_name", "")).border = thin_border
        ws.cell(row=row_idx, column=7, value=_format_changes(log.get("changes", []))).border = thin_border

        for col in range(1, 8):
            ws.cell(row=row_idx, column=col).alignment = data_alignment

    # 列宽
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 24
    ws.column_dimensions["G"].width = 50

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"操作日志_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@bp.route("/logs/delete", methods=["POST"])
def delete_logs():
    """删除选中日志"""
    log_ids = request.form.getlist("log_ids[]")

    if not log_ids:
        if _wants_json():
            return api_error("请选择要删除的日志", "NO_LOGS_SELECTED")
        flash("请选择要删除的日志", "error")
        return redirect(request.headers.get("Referer", "/"))

    service = _get_service()
    if not service:
        if _wants_json():
            return api_error("服务不可用", "SERVICE_UNAVAILABLE", 500)
        flash("服务不可用", "error")
        return redirect(request.headers.get("Referer", "/"))

    ids_to_delete = []
    for lid in log_ids:
        try:
            ids_to_delete.append(int(lid))
        except (ValueError, TypeError):
            pass

    try:
        count = service.delete_logs(ids_to_delete)
    except Exception as e:
        logger.exception("删除日志失败")
        if _wants_json():
            return api_error(f"删除日志失败: {e}", "DELETE_FAILED", 500)
        flash(f"删除日志失败: {e}", "error")
        return redirect(request.headers.get("Referer", "/"))

    if _wants_json():
        return api_success({"deleted_count": count}, f"成功删除 {count} 条日志")

    flash(f"成功删除 {count} 条日志", "success")
    return redirect("/card/logs")
