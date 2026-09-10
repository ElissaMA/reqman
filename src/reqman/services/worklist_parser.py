"""解析工作清单 Excel 文件

支持格式：
- 数据从 Row 7 开始
- 列 B = 工卡号, 列 E = 类型, 列 F = 专业
- 列 H = 工卡描述, 列 L = 备注
- Row 2 = 飞机信息
"""

import logging

import openpyxl

from .work_package_matcher import is_cancelled_item

logger = logging.getLogger(__name__)

# 文件大小上限
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
# 最大行数
MAX_ROWS = 5000


class WorklistError(Exception):
    """工作清单解析异常"""

    def __init__(self, message: str, detail: str = ""):
        self.message = message
        self.detail = detail
        super().__init__(message)


def parse_worklist(file_path: str, file_type: str = "例行") -> dict:
    """解析工作清单文件

    Args:
        file_path: Excel 文件路径
        file_type: '例行' 或 '其他'

    Returns:
        dict: {"aircraft_info": {...}, "items": [...], "total_count": N}

    Raises:
        WorklistError: 文件格式错误
    """
    # 文件大小检查
    import os
    size = os.path.getsize(file_path)
    if size > MAX_FILE_SIZE:
        raise WorklistError("文件过大", f"最大允许 10 MB，当前 {size / 1024 / 1024:.1f} MB")

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except (OSError, TypeError, ValueError) as e:
        raise WorklistError("无法打开文件，请确认是有效的 .xlsx 格式", str(e))

    try:
        if not wb.sheetnames:
            raise WorklistError("Excel 文件中没有工作表")

        ws = wb.active
        if ws.max_row < 7:
            raise WorklistError("文件内容不足，数据应从第 7 行开始")

        if ws.max_row > MAX_ROWS:
            raise WorklistError("文件行数过多", f"最大允许 {MAX_ROWS} 行")

        aircraft_info = _parse_aircraft_info(ws)
        items = _parse_items(ws, file_type)

    finally:
        wb.close()

    return {
        "aircraft_info": aircraft_info,
        "items": items,
        "total_count": len(items),
    }


def _parse_aircraft_info(ws) -> dict:
    """从 Row 2 提取飞机信息"""
    info = {"reg": "", "type": "", "package": "", "description": "", "level": "", "date": ""}

    try:
        for row in ws.iter_rows(min_row=2, max_row=2, values_only=False):
            for cell in row:
                if not cell.value:
                    continue
                val = str(cell.value).strip()
                if cell.column == 2:       # B2 = 机号
                    info["reg"] = val
                elif cell.column == 4:     # D2 = 机型
                    info["type"] = val
                elif cell.column == 8:     # H2 = 定检级别/描述（同源）
                    info["description"] = val
                    info["level"] = val
                elif cell.column == 10:    # J2 = 包号
                    info["package"] = val
    except (AttributeError, ValueError) as e:
        logger.warning(f"解析飞机信息时出错: {e}")

    # 从 Row 4 C4 提取日期
    try:
        for row in ws.iter_rows(min_row=4, max_row=4, values_only=False):
            for cell in row:
                if cell.column == 3 and cell.value:
                    raw = str(cell.value).strip()
                    date_part = raw.split()[0] if raw else ""
                    if date_part:
                        info["date"] = date_part.replace("-", ".")
                    break
    except (AttributeError, ValueError) as e:
        logger.warning(f"解析日期时出错: {e}")

    return info


def _parse_items(ws, file_type: str) -> list[dict]:
    """从 Row 7+ 提取工卡列表"""
    items = []
    seen_codes = set()
    error_rows = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=7, max_row=ws.max_row, values_only=False), start=7):
        try:
            # 列 B (index 1) = 工卡号
            if len(row) < 2:
                continue

            code_cell = row[1]
            if not code_cell or not code_cell.value:
                continue

            task_code = str(code_cell.value).strip()
            if not task_code:
                continue
            if task_code in seen_codes:
                continue
            seen_codes.add(task_code)

            # 列 F (index 5) = 专业
            category = ""
            if len(row) > 5 and row[5].value:
                category = str(row[5].value).strip()

            # 专业映射
            if category == "机身":
                category = "机体"

            # 列 E (index 4) = 类型
            task_type = ""
            if len(row) > 4 and row[4].value:
                task_type = str(row[4].value).strip()

            # 列 H (index 7) = 工卡描述
            task_name = ""
            if len(row) > 7 and row[7].value:
                task_name = str(row[7].value).strip()

            # 列 L (index 11) = 备注
            remark = ""
            if len(row) > 11 and row[11].value:
                remark = str(row[11].value).strip()

            item = {
                "task_code": task_code,
                "task_name": task_name,
                "category": category,
                "task_type": task_type,
                "remark": remark,
                "source": file_type,
            }
            item["cancelled"] = is_cancelled_item(item)
            items.append(item)

        except (ValueError, AttributeError, TypeError) as e:
            error_rows.append(f"Row {row_idx}: {e}")
            continue

    if error_rows:
        logger.warning(f"解析工作清单时有 {len(error_rows)} 行出错: {error_rows[:5]}")

    return items


def merge_aircraft_info(info_list: list[dict]) -> dict:
    """合并多个文件的飞机信息，优先使用非空值"""
    merged = {"reg": "", "type": "", "package": "", "description": "", "level": "", "date": ""}
    for info in info_list:
        if not isinstance(info, dict):
            continue
        for key, value in merged.items():
            if info.get(key) and not value:
                merged[key] = info[key]
    return merged
