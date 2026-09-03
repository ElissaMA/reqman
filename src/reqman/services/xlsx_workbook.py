"""需求单 Excel 读写 — 迁移自上级目录库存查询工具（scal）

读航材区（表头"定检专业\n（航材）"）与备用区（"备用航材需求"），
回填 G 列库存并标红（库存<需求）/标黄（需求≤库存<需求+2）。
原文件只读，输出内存副本（BytesIO），不写磁盘。
"""
from __future__ import annotations

import datetime
import io
import re
from pathlib import Path
from typing import NamedTuple

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

MATERIAL_HEADER = "定检专业\n（航材）"
SPARE_HEADER = "备用航材需求"

COL_PN = 3     # C
COL_NAME = 2   # B
COL_QTY = 5    # E
COL_STOCK = 7  # G


class DemandRow(NamedTuple):
    part_number: str
    name: str
    qty: float
    stock_cell: str
    row_idx: int


def _find_data_start(ws: openpyxl.Worksheet, header_keyword: str) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=1):
        cell = row[0]
        if cell.value and header_keyword in str(cell.value):
            for r in range(cell.row + 1, ws.max_row + 1):
                val = ws.cell(r, COL_PN).value
                if val and str(val).strip() not in ("", "件号"):
                    return r
    return None


def _next_section_row(ws: openpyxl.Worksheet, start: int) -> int | None:
    for r in range(start, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and re.match(r"^[三四五六七八九十]、", str(val).strip()):
            return r
    return None


def _parse_qty(raw: str | float | None) -> float:
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    m = re.search(r"[\d.]+", str(raw))
    return float(m.group()) if m else 0.0


def read_demand(path: str | Path) -> list[DemandRow]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    if ws is None:
        raise ValueError("No active sheet")

    rows: list[DemandRow] = []
    mat_start = _find_data_start(ws, MATERIAL_HEADER)
    spare_start = _find_data_start(ws, SPARE_HEADER)

    if mat_start:
        end = _next_section_row(ws, mat_start) or (ws.max_row + 1)
        if spare_start and spare_start > mat_start:
            end = min(end, spare_start)
        for r in range(mat_start, end):
            pn = ws.cell(r, COL_PN).value
            if not pn or not str(pn).strip():
                continue
            name = str(ws.cell(r, COL_NAME).value or "").strip()
            qty = _parse_qty(ws.cell(r, COL_QTY).value)
            rows.append(DemandRow(
                str(pn).strip().upper(), name, qty,
                get_column_letter(COL_STOCK) + str(r), r,
            ))

    if spare_start:
        end = _next_section_row(ws, spare_start) or (ws.max_row + 1)
        for r in range(spare_start, end):
            pn = ws.cell(r, COL_PN).value
            if not pn or not str(pn).strip():
                continue
            name = str(ws.cell(r, COL_NAME).value or "").strip()
            qty = _parse_qty(ws.cell(r, COL_QTY).value)
            rows.append(DemandRow(
                str(pn).strip().upper(), name, qty,
                get_column_letter(COL_STOCK) + str(r), r,
            ))

    wb.close()
    return rows


def write_inventory_copy(
    source_path: str | Path,
    results: dict[str, float],
    pn_qty: dict[str, list[float]],
    pn_cells: dict[str, list[str]],
    timestamp_suffix: str | None = None,
    output_stem: str | None = None,
    warning_thresholds: dict | None = None,
) -> tuple[io.BytesIO, str]:
    source = Path(source_path)
    if timestamp_suffix is None:
        timestamp_suffix = datetime.datetime.now(datetime.timezone.utc).astimezone().strftime("%Y%m%d_%H%M%S")

    stem = output_stem or source.stem
    filename = f"{stem}_库存已填_{timestamp_suffix}.xlsx"

    wb = openpyxl.load_workbook(source)
    ws = wb.active
    if ws is None:
        raise ValueError("No active sheet")

    for pn, cells in pn_cells.items():
        stock_val = results.get(pn, 0.0)
        qtys = pn_qty.get(pn, [0.0])

        for cell_addr, qty in zip(cells, qtys):
            cell = ws[cell_addr]
            cell.value = int(stock_val) if stock_val == int(stock_val) else stock_val

            if stock_val < qty:
                fill = RED_FILL
            else:
                thr = warning_thresholds.get(pn) if warning_thresholds else None
                if thr is not None:
                    fill = YELLOW_FILL if stock_val < thr else None
                else:
                    fill = YELLOW_FILL if qty <= stock_val < qty + 2 else None

            if fill:
                row = cell.row
                for c in range(2, 8):  # B(2) ~ G(7)
                    ws.cell(row, c).fill = fill

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf, filename
