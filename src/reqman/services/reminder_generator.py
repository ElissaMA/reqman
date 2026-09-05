"""提醒单生成器 — 以空白模板生成《定检工作提醒单》

填充规则：
- 行1 标题（模板保留）；行2-4 飞机信息
- 行7+ 数据区：按专业列写入（电子→A、发动机→B、机体→C）工卡名称
- 例行清单 → 黑字；其他清单 → 红字；重点提醒 → 黄底 RGB(255,255,0)
- 未识别工卡不输出（由预览页 new_cards 处理）
"""
from __future__ import annotations

import io
import os
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl.styles import Font, PatternFill

from ..config import REMINDER_TEMPLATE_FILE
from ..utils.template_cache import load_template

COL_MAP = {"电子": 1, "发动机": 2, "机体": 3}   # A / B / C
DATA_START = 7
BLACK = "FF000000"
RED = "FFFF0000"
YELLOW_FILL = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")


def generate_reminder(form_data: dict, items: list[dict]) -> tuple[io.BytesIO, str]:
    if not os.path.exists(REMINDER_TEMPLATE_FILE):
        raise FileNotFoundError(f"模板文件不存在：{REMINDER_TEMPLATE_FILE}")

    reg = form_data.get("reg", "XXXX").removeprefix("B-")
    date_str = form_data.get("date", datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d"))
    desc = form_data.get("description", "")
    filename = f"定检工作提醒单（B-{reg} {desc}）{date_str}.xlsx"

    wb = load_template(REMINDER_TEMPLATE_FILE)
    ws = wb["工卡提醒"]

    ws["A2"].value = f"时间：{date_str}"
    ws["B2"].value = f"定检级别：{form_data.get('level', '')}"
    ws["C2"].value = f"总份数：{form_data.get('routine_count', '')}+{form_data.get('other_count', '')}"
    ws["A3"].value = f"机号：{form_data.get('reg', '')}"
    ws["B3"].value = f"机型：{form_data.get('aircraft_type', '')}"
    ws["C3"].value = f"APU型号：{form_data.get('apu', '')}"
    ws["A4"].value = f"FSN：{form_data.get('fsn', '')}"
    ws["B4"].value = f"MSN：{form_data.get('msn', '')}"
    ws["C4"].value = "上次定检："

    for item in items:
        col = COL_MAP.get(item.get("category", ""))
        if col is None:
            continue
        row = _next_row(ws, col)
        cell = ws.cell(row=row, column=col)
        cell.value = item.get("task_name", "")
        is_other = item.get("source") == "其他"
        cell.font = Font(name="宋体", size=11, color=RED if is_other else BLACK)
        if item.get("reminder_type") == "重点提醒":
            cell.fill = YELLOW_FILL

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    buffer.seek(0)
    return buffer, filename


def _next_row(ws, col: int) -> int:
    """找到该列数据区下一个空行（从 DATA_START 起）。"""
    row = DATA_START
    while ws.cell(row=row, column=col).value is not None and ws.cell(row=row, column=col).value != "":
        row += 1
    return row