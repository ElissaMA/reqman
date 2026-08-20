"""VBA 配置文件迁移 — 一次性脚本（scripts/import_vba_config.py 调用）

规则：
- 电子/发动机/机体提醒 → 已存在卡设 reminder_type="一般提醒"；不存在→弃用
- 重点工卡 → 已存在卡设 reminder_type="重点提醒"（覆盖一般）；不存在→弃用
- 无需提醒表不迁移
- 导入卡初始 card_ok=False（需人工确认）
- 飞机信息按 reg 比对防重复：存在→更新，不存在→导入
- 弃用清单写入 output/vba_discard.txt
"""
from __future__ import annotations

from pathlib import Path
from typing import NamedTuple

import openpyxl

from ..models.json_store import JsonStore

GENERAL_SHEETS = {"电子提醒": "电子", "发动机提醒": "发动机", "机体提醒": "机体"}
KEY_SHEET = "重点工卡"
AIRCRAFT_SHEET = "飞机信息"


class ImportResult(NamedTuple):
    general: int
    key: int
    discarded: list[str]
    aircraft_added: int
    aircraft_updated: int


def import_vba_config(config_xlsx: str, store: JsonStore) -> ImportResult:
    wb = openpyxl.load_workbook(config_xlsx, read_only=True, data_only=True)
    discarded: list[str] = []
    general = key = aircraft_added = aircraft_updated = 0

    # 1) 专业提醒：一般提醒（仅匹配已存在工卡）
    for sheet_name in GENERAL_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = str(row[0]).strip() if row and row[0] else ""
            if not code:
                continue
            card = store.find_by_code(code)
            if card is None:
                discarded.append(f"[{sheet_name}] {code}")
                continue
            store.update(card["id"], reminder_type="一般提醒", card_ok=False)
            general += 1

    # 2) 重点工卡：重点提醒（覆盖一般）
    if KEY_SHEET in wb.sheetnames:
        ws = wb[KEY_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = str(row[0]).strip() if row and row[0] else ""
            if not code:
                continue
            card = store.find_by_code(code)
            if card is None:
                discarded.append(f"[{KEY_SHEET}] {code}")
                continue
            store.update(card["id"], reminder_type="重点提醒", card_ok=False)
            key += 1

    # 3) 飞机信息：reg 比对防重复
    if AIRCRAFT_SHEET in wb.sheetnames:
        ws = wb[AIRCRAFT_SHEET]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            reg = str(row[0]).strip()
            model = str(row[1] or "").strip()
            fsn = str(row[2] or "").strip()
            msn = str(row[3] or "").strip()
            apu = str(row[4] or "").strip()
            ac = store.find_aircraft_by_reg(reg)
            if ac is None:
                store.add_aircraft(reg=reg, model=model, fsn=fsn, msn=msn, apu=apu)
                aircraft_added += 1
            else:
                store.update_aircraft(ac["id"], model=model, fsn=fsn, msn=msn, apu=apu)
                aircraft_updated += 1

    wb.close()

    # 4) 弃用清单落盘
    if discarded:
        out = Path("output")
        out.mkdir(exist_ok=True)
        (out / "vba_discard.txt").write_text(
            "\n".join(discarded) + "\n", encoding="utf-8")

    return ImportResult(general, key, discarded, aircraft_added, aircraft_updated)